import hashlib
import io
import re
from datetime import datetime, timezone
import subprocess
import sys
import tempfile
import types
import zipfile
from pathlib import Path
from unittest import mock
from urllib.parse import unquote

import numpy as np
import pymupdf
import pytest
from PIL import Image, ImageChops, ImageDraw, ImageFilter

import app as euamopdf


@pytest.fixture
def client():
    return euamopdf.app.test_client()


def noise(size, sigma, seed=0):
    """Ruído gaussiano como o do Image.effect_noise, mas sempre o mesmo: teste reproduzível."""
    values = np.random.default_rng(seed).normal(128, sigma, (size[1], size[0]))
    return Image.fromarray(values.clip(0, 255).astype(np.uint8))


def make_pdf(pages=3, **save_options):
    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page().insert_text((72, 72), f"Pagina {i + 1}")
    return doc.tobytes(**save_options)


def make_image(mode="RGB", size=(200, 100), color=(0, 0, 255), fmt="PNG", **save_options):
    buf = io.BytesIO()
    Image.new(mode, size, color).save(buf, fmt, **save_options)
    return buf.getvalue()


def post(client, action, *files, headers=None, **form):
    """files: pares (nome, conteúdo)."""
    data = {"action": action, "file": [(io.BytesIO(content), name) for name, content in files], **form}
    return client.post("/convert", data=data, content_type="multipart/form-data", headers=headers or {})


def texts(pdf):
    return [page.get_text().strip() for page in pymupdf.open(stream=pdf, filetype="pdf")]


def unzip(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    return {name: z.read(name) for name in z.namelist()}


# --- Organizar ---

def test_merge_keeps_upload_order(client):
    r = post(client, "merge-pdf", ("b.pdf", make_pdf(2)), ("a.pdf", make_pdf(1)))
    assert r.status_code == 200
    assert texts(r.data) == ["Pagina 1", "Pagina 2", "Pagina 1"]


def test_split_every_page_into_zip(client):
    r = post(client, "split-pdf", ("doc.pdf", make_pdf(3)))
    files = unzip(r.data)
    assert sorted(files) == ["pag_1.pdf", "pag_2.pdf", "pag_3.pdf"]
    assert texts(files["pag_2.pdf"]) == ["Pagina 2"]


@pytest.mark.parametrize("spec, expected", [
    ("2-3", ["Pagina 2", "Pagina 3"]),
    ("3, 1", ["Pagina 3", "Pagina 1"]),
    ("1-2,3", ["Pagina 1", "Pagina 2", "Pagina 3"]),
])
def test_split_page_list(client, spec, expected):
    r = post(client, "split-pdf", ("doc.pdf", make_pdf(3)), pages=spec)
    assert r.status_code == 200
    assert texts(r.data) == expected


@pytest.mark.parametrize("spec", ["0-1", "3-1", "1-9", "4", "-2", "abc", " , "])
def test_split_rejects_invalid_pages(client, spec):
    r = post(client, "split-pdf", ("doc.pdf", make_pdf(3)), pages=spec)
    assert r.status_code == 400


def make_photo_pdf():
    """PDF com uma foto de alta resolução numa área pequena da página."""
    buf = io.BytesIO()
    noise((1200, 1200), 60).convert("RGB").save(buf, "PNG")
    doc = pymupdf.open()
    doc.new_page().insert_image(pymupdf.Rect(72, 72, 272, 272), stream=buf.getvalue())
    return doc.tobytes()


@pytest.mark.parametrize("level", ["recomendada", "forte"])
def test_compress_reduces_images(client, level):
    original = make_photo_pdf()
    r = post(client, "compress-pdf", ("foto.pdf", original), level=level)
    assert r.status_code == 200
    assert len(r.data) < len(original) / 2
    assert "Reduzido de" in unquote(r.headers["X-Mensagem"])


def make_scan_pdf(dpi):
    """Página A4 escaneada: uma foto ocupando a página inteira na resolução dada."""
    size = (round(595 / 72 * dpi), round(842 / 72 * dpi))
    doc = pymupdf.open()
    doc.new_page().insert_image(pymupdf.paper_rect("a4"), stream=make_photo(size, quality=90))
    return doc.tobytes()


@pytest.mark.parametrize("level, width", [("recomendada", 1240), ("forte", 794)])
def test_compress_pdf_reaches_the_target_resolution(client, level, width):
    # 200 DPI -> 150 ou 96 DPI: fatores que não são potência de 2, que o MuPDF não reduzia
    r = post(client, "compress-pdf", ("scan.pdf", make_scan_pdf(200)), level=level)
    [image] = pymupdf.open(stream=r.data, filetype="pdf")[0].get_image_info()
    assert abs(image["width"] - width) <= 2
    assert image["bbox"][2] == pytest.approx(595, abs=1)  # continua ocupando a página


def test_compress_pdf_keeps_transparent_and_small_lossless_images(client):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(0, 0, 200, 200), stream=make_photo((1000, 1000), fmt="PNG", mode="RGBA"))
    page.insert_image(pymupdf.Rect(0, 300, 300, 600), stream=make_photo((300, 300), fmt="PNG"))  # 72 DPI
    r = post(client, "compress-pdf", ("doc.pdf", doc.tobytes()))
    page = pymupdf.open(stream=r.data, filetype="pdf")[0]
    assert sorted(info["width"] for info in page.get_image_info()) == [300, 1000]
    assert [item[1] > 0 for item in page.get_images(full=True)].count(True) == 1  # a transparência continua


def test_compress_never_makes_file_bigger(client):
    original = make_pdf(1, garbage=4, deflate=True, clean=True, use_objstms=1)
    r = post(client, "compress-pdf", ("texto.pdf", original))
    assert len(r.data) <= len(original)
    assert "já está otimizado" in unquote(r.headers["X-Mensagem"])


def rotations(pdf):
    return [page.rotation for page in pymupdf.open(stream=pdf, filetype="pdf")]


def test_rotate_all_pages(client):
    r = post(client, "rotate-pdf", ("doc.pdf", make_pdf(3)), angle="90")
    assert rotations(r.data) == [90, 90, 90]


def test_rotate_selected_pages_adds_to_current_rotation(client):
    once = post(client, "rotate-pdf", ("doc.pdf", make_pdf(3)), angle="270", pages="1, 3").data
    twice = post(client, "rotate-pdf", ("doc.pdf", once), angle="180", pages="1").data
    assert rotations(twice) == [90, 0, 270]


def test_rotate_requires_valid_angle(client):
    assert post(client, "rotate-pdf", ("doc.pdf", make_pdf()), angle="45").status_code == 400


def test_protect_then_unlock(client):
    r = post(client, "protect-pdf", ("doc.pdf", make_pdf(2)), password="ação123", password2="ação123")
    locked = pymupdf.open(stream=r.data, filetype="pdf")
    assert locked.needs_pass
    assert locked.authenticate("ação123")

    assert post(client, "unlock-pdf", ("doc.pdf", r.data), password="errada").get_data(as_text=True) == "Senha incorreta."
    r = post(client, "unlock-pdf", ("doc.pdf", r.data), password="ação123")
    assert not pymupdf.open(stream=r.data, filetype="pdf").needs_pass
    assert texts(r.data) == ["Pagina 1", "Pagina 2"]


def test_protect_requires_matching_passwords(client):
    assert post(client, "protect-pdf", ("doc.pdf", make_pdf()), password="", password2="").status_code == 400
    assert post(client, "protect-pdf", ("doc.pdf", make_pdf()), password="a", password2="b").status_code == 400


def test_unlock_removes_restrictions_without_password(client):
    restricted = make_pdf(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="dono",
                          permissions=pymupdf.PDF_PERM_ACCESSIBILITY)
    r = post(client, "unlock-pdf", ("doc.pdf", restricted))
    assert r.status_code == 200
    assert pymupdf.open(stream=r.data, filetype="pdf").metadata["encryption"] is None


def test_unlock_plain_pdf_is_reported(client):
    r = post(client, "unlock-pdf", ("doc.pdf", make_pdf()))
    assert r.status_code == 400


def visual_lines(page):
    """Linhas de texto com direção e posição como o usuário vê a página (já girada)."""
    m = page.rotation_matrix
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            direction = pymupdf.Point(line["dir"]) * m - pymupdf.Point(0, 0) * m
            text = "".join(span["text"] for span in line["spans"])
            yield text, (round(direction.x, 2), round(direction.y, 2)), pymupdf.Rect(line["bbox"]) * m


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_watermark_rises_diagonally_on_any_rotation(client, rotation):
    doc = pymupdf.open()
    doc.new_page().set_rotation(rotation)
    r = post(client, "watermark-pdf", ("doc.pdf", doc.tobytes()), text="CONFIDENCIAL")
    page = pymupdf.open(stream=r.data, filetype="pdf")[0]
    [(text, direction, bbox)] = visual_lines(page)
    assert text == "CONFIDENCIAL"
    assert direction == (0.71, -0.71)  # sobe da esquerda para a direita
    assert abs(bbox.x0 + bbox.x1 - page.rect.width) < 20  # centralizada


@pytest.mark.parametrize("text", ["", "   ", "Aprovado ✅"])
def test_watermark_rejects_invalid_text(client, text):
    assert post(client, "watermark-pdf", ("doc.pdf", make_pdf()), text=text).status_code == 400


@pytest.mark.parametrize("fmt, expected", [("n", "2"), ("n-de-t", "2 / 3"), ("pagina", "Página 2 de 3")])
def test_number_pages_formats(client, fmt, expected):
    r = post(client, "number-pages", ("doc.pdf", make_pdf(3)), format=fmt)
    page = pymupdf.open(stream=r.data, filetype="pdf")[1]
    assert expected in [text for text, _, _ in visual_lines(page)]


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_number_pages_at_bottom_center_on_any_rotation(client, rotation):
    doc = pymupdf.open()
    doc.new_page().set_rotation(rotation)
    r = post(client, "number-pages", ("doc.pdf", doc.tobytes()))
    page = pymupdf.open(stream=r.data, filetype="pdf")[0]
    [(text, direction, bbox)] = visual_lines(page)
    assert (text, direction) == ("1", (1.0, 0.0))  # horizontal, como o usuário lê
    assert bbox.y1 > page.rect.height - 30
    assert abs(bbox.x0 + bbox.x1 - page.rect.width) < 5


def make_scanned_pdf(text, rotation=0):
    """PDF só com a imagem do texto, como um escaneamento. Com rotação, a imagem é gravada
    girada e a página a exibe em pé, como fazem muitos scanners."""
    source = pymupdf.open()
    source.new_page().insert_text((72, 100), text, fontsize=18)
    pix = source[0].get_pixmap(dpi=150)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).rotate(rotation, expand=True)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = pymupdf.open()
    page = doc.new_page(width=img.width * 72 / 150, height=img.height * 72 / 150)
    page.insert_image(page.rect, stream=buf.getvalue())
    page.set_rotation(rotation)
    return doc.tobytes()


needs_tesseract = pytest.mark.skipif(
    not (euamopdf.find_tessdata() and Path(euamopdf.find_tessdata(), "por.traineddata").exists()),
    reason="Tesseract com português não instalado")


@needs_tesseract
@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_ocr_makes_scanned_text_searchable(client, rotation):
    scanned = make_scanned_pdf("Contrato de locação número 42", rotation)
    assert texts(scanned) == [""]
    r = post(client, "ocr-pdf", ("scan.pdf", scanned))
    assert r.status_code == 200
    page = pymupdf.open(stream=r.data, filetype="pdf")[0]
    assert " ".join(page.get_text().split()) == "Contrato de locação número 42"
    assert page.search_for("locação")  # dá para buscar


def test_invisible_text_ignores_degenerate_ocr_boxes():
    page = pymupdf.open().new_page()
    euamopdf.add_invisible_text(page, [(10, 10, 10, 20, "vazia", 0, 0, 0), (10, 20, 50, 20, "achatada", 0, 1, 0),
                                       (10, 30, 60, 45, "certa", 0, 2, 0)])
    assert page.get_text().split() == ["certa"]


@needs_tesseract
def test_ocr_skips_pages_that_already_have_text(client):
    r = post(client, "ocr-pdf", ("doc.pdf", make_pdf()))
    assert r.status_code == 400
    assert "já têm texto" in r.get_data(as_text=True)


def test_ocr_without_tesseract_explains_what_to_install(client, monkeypatch):
    monkeypatch.setattr(euamopdf, "find_tessdata", lambda: None)
    r = post(client, "ocr-pdf", ("scan.pdf", make_pdf()))
    assert r.status_code == 400
    assert "Tesseract" in r.get_data(as_text=True)


def test_organize_reorders_rotates_duplicates_and_deletes(client):
    order = '[{"pagina": 3, "giro": 90}, {"pagina": 1}, {"pagina": 1, "giro": -90}]'
    r = post(client, "organize-pdf", ("doc.pdf", make_pdf(3)), order=order)
    assert texts(r.data) == ["Pagina 3", "Pagina 1", "Pagina 1"]
    assert rotations(r.data) == [90, 0, 270]


@pytest.mark.parametrize("order", ["", "[]", "nao e json", "[1, 2]", '[{"pagina": 4}]', '[{"pagina": 0}]',
                                   '[{"pagina": 1, "giro": 45}]', '{"pagina": 1}', '[{"page": 1}]'])
def test_organize_rejects_invalid_order(client, order):
    assert post(client, "organize-pdf", ("doc.pdf", make_pdf(3)), order=order).status_code == 400


def test_page_thumbnails(client):
    r = client.post("/pages", data={"file": (io.BytesIO(make_pdf(2)), "a.pdf"), "miniaturas": "1"},
                    content_type="multipart/form-data")
    thumbs = r.get_json()["miniaturas"]
    assert len(thumbs) == 2
    assert all(t.startswith("data:image/jpeg;base64,") for t in thumbs)


def make_photo(size=(1600, 1200), fmt="JPEG", quality=95, exif=None, mode="RGB"):
    """Imagem com gradientes e ruído, que se comprime como uma foto de verdade."""
    w, h = size
    img = Image.merge("RGB", [Image.linear_gradient("L").resize(size), Image.radial_gradient("L").resize(size),
                              noise(size, 30)])
    if mode == "RGBA":
        img.putalpha(Image.linear_gradient("L").resize(size))  # topo transparente, base opaca
    buf = io.BytesIO()
    options = {"quality": quality} if fmt in ("JPEG", "WEBP") else {}
    img.save(buf, fmt, **options, **({"exif": exif} if exif else {}))
    return buf.getvalue()


def opened(data):
    return Image.open(io.BytesIO(data))


def compress_image(client, *files, **form):
    return post(client, "compress-image", *files, **form)


def test_compress_image_levels_get_progressively_smaller(client):
    photo = make_photo()
    sizes = [len(compress_image(client, ("foto.jpg", photo), image_level=level).data)
             for level in ("leve", "recomendada", "forte", "extrema")]
    assert sizes == sorted(sizes, reverse=True)
    assert sizes[0] < len(photo)
    r = compress_image(client, ("foto.jpg", photo), image_level="recomendada")
    assert r.headers["Content-Disposition"].endswith("foto_comprimida.jpg")
    assert opened(r.data).format == "JPEG"
    assert "Reduzida de" in unquote(r.headers["X-Mensagem"])


def test_compress_image_resizes_only_down(client):
    r = compress_image(client, ("foto.jpg", make_photo((3000, 2000))), max_size="1280")
    assert opened(r.data).size == (1280, 853)
    r = compress_image(client, ("pequena.jpg", make_photo((600, 400))), max_size="1280")
    assert opened(r.data).size == (600, 400)


@pytest.mark.parametrize("target, fmt, ext", [("jpeg", "JPEG", ".jpg"), ("webp", "WEBP", ".webp"), ("png", "PNG", ".png")])
def test_compress_image_converts_format(client, target, fmt, ext):
    r = compress_image(client, ("foto.png", make_photo((600, 400), fmt="PNG")), image_format=target)
    assert opened(r.data).format == fmt
    assert r.headers["Content-Disposition"].endswith(f"foto_comprimida{ext}")


def test_compress_image_png_to_jpg_puts_transparency_on_white(client):
    png = make_photo((200, 200), fmt="PNG", mode="RGBA")
    r = compress_image(client, ("logo.png", png), image_format="jpeg")
    assert opened(r.data).getpixel((100, 0))[0] > 240  # topo transparente vira branco


def test_compress_image_png_keeps_transparency_and_reduces_colors(client):
    png = make_photo((400, 300), fmt="PNG", mode="RGBA")
    r = compress_image(client, ("logo.png", png), image_level="forte")
    img = opened(r.data)
    assert img.format == "PNG" and img.mode == "P"
    assert len(img.getcolors(256)) <= 128
    assert "transparency" in img.info or img.convert("RGBA").getextrema()[3][0] < 255
    assert len(r.data) < len(png)


def exif_with_location():
    exif = Image.Exif()
    exif[0x010F] = "Camera do celular"   # fabricante
    exif[0x0112] = 6                      # orientação: gravada deitada
    exif[0x8825] = {2: (23.0, 32.0, 0.0)}   # GPS: latitude (via get_ifd, o Pillow < 11.1 não grava)
    return exif


def test_compress_image_removes_photo_data_by_default_in_the_ui(client):
    photo = make_photo((400, 300), exif=exif_with_location())
    r = compress_image(client, ("foto.jpg", photo), strip_metadata="1")
    img = opened(r.data)
    assert not img.getexif()
    assert img.size == (300, 400)  # a orientação foi aplicada antes de tirar os dados


def test_compress_image_can_keep_photo_data(client):
    photo = make_photo((400, 300), exif=exif_with_location())
    r = compress_image(client, ("foto.jpg", photo))
    exif = opened(r.data).getexif()
    assert exif[0x010F] == "Camera do celular"
    assert 0x0112 not in exif  # orientação já aplicada: não pode ser girada de novo
    assert exif.get_ifd(0x8825)


def test_compress_image_keeps_already_compressed_original(client):
    small = make_photo((400, 300), quality=30)
    r = compress_image(client, ("leve.jpg", small), image_level="leve")
    assert r.data == small
    assert "já estava bem comprimida" in unquote(r.headers["X-Mensagem"])


def test_compress_image_says_when_the_result_is_bigger(client):
    photo = make_photo((400, 300))  # foto JPG que, em PNG sem perda, fica bem maior
    r = compress_image(client, ("foto.jpg", photo), image_level="leve", image_format="png")
    message = unquote(r.headers["X-Mensagem"])
    assert len(r.data) > len(photo)
    assert message.startswith("Ficou maior") and "% maior" in message


def test_compress_many_images_into_zip_without_name_clashes(client):
    r = compress_image(client, ("foto.jpg", make_photo((400, 300))), ("foto.jpg", make_photo((300, 200))),
                       ("tela.png", make_photo((300, 200), fmt="PNG")))
    assert r.headers["Content-Disposition"].endswith("Imagens_comprimidas.zip")
    assert sorted(unzip(r.data)) == ["foto_comprimida.jpg", "foto_comprimida_2.jpg", "tela_comprimida.png"]
    assert unquote(r.headers["X-Mensagem"]).startswith("3 imagens: de")


@pytest.mark.parametrize("name, content", [("anim.gif", make_image(fmt="GIF")), ("foto.jpg", b"nao e imagem")])
def test_compress_image_rejects_other_files(client, name, content):
    assert compress_image(client, (name, content)).status_code == 400


@pytest.mark.filterwarnings("ignore::PIL.Image.DecompressionBombWarning")  # o app o silencia; o pytest reativa
def test_200_megapixel_photos_are_accepted(client):
    """O Pillow recusa por padrão imagens acima de ~179 MP, e a foto de um celular de 200 MP
    (16320 x 12240) era dada como "não é uma imagem válida" em todas as ferramentas de imagem."""
    buf = io.BytesIO()
    Image.new("L", (16320, 12240), 128).save(buf, "JPEG", quality=50)
    photo = buf.getvalue()
    r = compress_image(client, ("foto.jpg", photo), max_size="1920")
    assert r.status_code == 200 and max(opened(r.data).size) == 1920
    r = client.post("/metadata", data={"file": (io.BytesIO(photo), "foto.jpg")}, content_type="multipart/form-data")
    assert r.status_code == 200


def test_image_above_the_limit_gets_its_own_message(client, monkeypatch):
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 1000)  # o limite de verdade é de 300 MP
    for action in ("compress-image", "jpg-to-pdf", "edit-metadata"):
        r = post(client, action, ("foto.png", make_image()))
        assert r.status_code == 400 and "grande demais" in r.get_data(as_text=True), action


def estimate(client, action, *files, **form):
    r = client.post("/estimate", data={"action": action, "file": [(io.BytesIO(c), n) for n, c in files], **form},
                    content_type="multipart/form-data")
    return r


@pytest.mark.parametrize("form", [
    {"image_level": "recomendada"},
    {"image_level": "extrema", "image_format": "webp"},
    {"image_level": "forte", "image_format": "png", "max_size": "1280"},
    {"image_level": "leve", "max_size": "800"},
])
def test_image_estimate_is_close_to_the_real_result(client, form):
    photos = [("a.jpg", make_photo((3000, 2000))), ("b.jpg", make_photo((1600, 1200), quality=80))]  # a primeira passa pelo mosaico
    r = estimate(client, "compress-image", *photos, **form)
    assert r.status_code == 200
    body = r.get_json()
    assert body["antes"] == sum(len(content) for _, content in photos)
    real = sum(len(data) for data in unzip(compress_image(client, *photos, **form).data).values())
    assert body["depois"] == pytest.approx(real, rel=0.2)


def test_image_estimate_samples_at_least_three_images(client, monkeypatch):
    monkeypatch.setattr(euamopdf, "ESTIMATE_SECONDS", 0)  # prazo esgotado já na primeira imagem
    files = [("foto.jpg", make_photo((1600, 1200))), ("logo.png", make_photo((600, 600), fmt="PNG", mode="RGBA"))]
    form = {"image_level": "forte", "image_format": "webp"}
    body = estimate(client, "compress-image", *files, **form).get_json()
    real = sum(len(data) for data in unzip(compress_image(client, *files, **form).data).values())
    assert body["depois"] == pytest.approx(real, rel=0.15)


def test_image_estimate_never_exceeds_an_original_that_stays(client):
    small = make_photo((400, 300), quality=30)
    body = estimate(client, "compress-image", ("leve.jpg", small), image_level="leve").get_json()
    assert body["depois"] == len(small)


@pytest.mark.parametrize("level", ["recomendada", "forte"])
def test_pdf_estimate_is_close_to_the_real_result(client, level):
    doc = pymupdf.open()
    for dpi in (200, 250, 150):  # três páginas escaneadas, cada uma numa resolução
        doc.insert_pdf(pymupdf.open(stream=make_scan_pdf(dpi), filetype="pdf"))
    pdf = doc.tobytes()
    body = estimate(client, "compress-pdf", ("scan.pdf", pdf), level=level).get_json()
    real = len(post(client, "compress-pdf", ("scan.pdf", pdf), level=level).data)
    assert body["antes"] == len(pdf)
    assert body["depois"] == pytest.approx(real, rel=0.1)


def test_pdf_estimate_with_uncompressed_image(client):
    # Imagem gravada sem compressão: ao salvar, o PDF a comprime, e a estimativa precisa contar com isso
    doc = pymupdf.open()
    page = doc.new_page()
    xref = page.insert_image(page.rect, stream=make_photo((1240, 1754), fmt="PNG"))
    doc.update_stream(xref, pymupdf.Pixmap(doc, xref).samples, compress=False)
    doc.xref_set_key(xref, "Filter", "null")
    doc.xref_set_key(xref, "DecodeParms", "null")
    pdf = doc.tobytes()
    body = estimate(client, "compress-pdf", ("scan.pdf", pdf), level="forte").get_json()
    real = len(post(client, "compress-pdf", ("scan.pdf", pdf), level="forte").data)
    assert 0 < body["depois"] < body["antes"]
    assert body["depois"] == pytest.approx(real, rel=0.2)


def test_pdf_estimate_never_exceeds_the_original(client):
    pdf = make_pdf(1, garbage=4, deflate=True, clean=True, use_objstms=1)
    body = estimate(client, "compress-pdf", ("texto.pdf", pdf)).get_json()
    assert body["depois"] <= body["antes"]


@pytest.mark.parametrize("action, files, status", [
    ("merge-pdf", [("a.pdf", make_pdf())], 400),                       # ferramenta sem estimativa
    ("compress-pdf", [("a.pdf", b"lixo")], 400),                       # arquivo inválido
    ("compress-pdf", [("a.pdf", make_pdf(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="1", owner_pw="1"))], 400),
    ("compress-image", [("a.gif", make_image(fmt="GIF"))], 400),       # tipo não aceito
])
def test_estimate_rejects_what_the_tool_would_reject(client, action, files, status):
    assert estimate(client, action, *files).status_code == status


def test_spread_picks_items_across_the_list():
    assert euamopdf.spread(list(range(10)), 3) == [0, 3, 6]
    assert euamopdf.spread([1, 2], 3) == [1, 2]


# --- Remover fundo ---

class FakeModel:
    """Faz as vezes do ISNet: a máscara é um disco no centro, qualquer que seja a imagem."""
    def __init__(self, found=True):
        yy, xx = np.mgrid[:1024, :1024]
        self.mask = ((yy - 512) ** 2 + (xx - 512) ** 2 < 300 ** 2).astype(np.float32) * found

    def get_inputs(self):
        return [types.SimpleNamespace(name="entrada")]

    def run(self, outputs, feeds):
        x = feeds["entrada"]
        assert x.shape == (1, 3, 1024, 1024) and x.dtype == np.float32 and -0.51 < x.min() and x.max() < 0.51
        return [self.mask[None, None]]


@pytest.fixture
def fake_model(monkeypatch):
    monkeypatch.setattr(euamopdf, "_bg_session", FakeModel())


def remove_bg(client, *files, **form):
    return post(client, "remove-background", *files, **form)


def test_remove_background_keeps_original_pixels_and_size(client, fake_model):
    photo = make_photo((1600, 1200), fmt="PNG")
    r = remove_bg(client, ("foto.png", photo))
    assert r.headers["Content-Disposition"].endswith("foto_sem_fundo.png")
    out, src = opened(r.data), opened(photo).convert("RGB")
    assert out.format == "PNG" and out.mode == "RGBA" and out.size == src.size
    alpha = np.asarray(out)[..., 3]
    assert alpha[600, 800] == 255 and alpha[0, 0] == 0
    inside = alpha == 255  # no objeto, os pixels são exatamente os da fonte
    assert np.array_equal(np.asarray(out)[..., :3][inside], np.asarray(src)[inside])


def test_remove_background_erases_what_was_behind_transparent_pixels(client, fake_model):
    r = remove_bg(client, ("foto.png", make_photo((400, 300), fmt="PNG")))
    pixels = np.asarray(opened(r.data))
    assert not pixels[pixels[..., 3] == 0].any()  # o fundo não fica escondido atrás da transparência


def test_remove_background_follows_photo_orientation_and_keeps_color_profile(client, fake_model):
    img = Image.open(io.BytesIO(make_photo((600, 400))))
    exif = Image.Exif()
    exif[0x0112] = 6
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif, icc_profile=b"perfil de cor de teste")
    out = opened(remove_bg(client, ("foto.jpg", buf.getvalue())).data)
    assert out.size == (400, 600)
    assert out.info.get("icc_profile") == b"perfil de cor de teste"


@pytest.mark.parametrize("form, color", [
    ({"background": "branco"}, (255, 255, 255)),
    ({"background": "preto"}, (0, 0, 0)),
    ({"background": "cor", "background_color": "#2F80ED"}, (47, 128, 237)),
])
def test_remove_background_paints_new_background(client, fake_model, form, color):
    photo = make_photo((400, 300), fmt="PNG")
    out = opened(remove_bg(client, ("foto.png", photo), **form).data)
    assert out.mode == "RGB"
    assert out.getpixel((0, 0)) == color
    assert out.getpixel((200, 150)) == opened(photo).convert("RGB").getpixel((200, 150))


def test_remove_background_webp_is_lossless(client, fake_model):
    photo = make_photo((400, 300), fmt="PNG")
    r = remove_bg(client, ("foto.png", photo), bg_format="webp")
    out = opened(r.data)
    assert out.format == "WEBP" and r.headers["Content-Disposition"].endswith(".webp")
    assert out.convert("RGB").getpixel((200, 150)) == opened(photo).convert("RGB").getpixel((200, 150))


def test_remove_background_keeps_existing_transparency(client, fake_model):
    png = make_photo((400, 400), fmt="PNG", mode="RGBA")  # transparência em degradê: topo transparente
    alpha = np.asarray(opened(remove_bg(client, ("logo.png", png)).data))[..., 3].astype(int)
    source = np.asarray(opened(png))[..., 3].astype(int)
    for y in (120, 200, 300):  # dentro do disco (raio ~117 px), vale a transparência original
        assert abs(alpha[y, 200] - source[y, 200]) <= 1
    assert alpha[200, 10] == 0  # fora do disco, transparente


def test_remove_background_many_images_and_no_subject(client, monkeypatch):
    monkeypatch.setattr(euamopdf, "_bg_session", FakeModel(found=False))
    r = remove_bg(client, ("a.jpg", make_photo((300, 200))), ("a.jpg", make_photo((200, 300))))
    assert sorted(unzip(r.data)) == ["a_sem_fundo.png", "a_sem_fundo_2.png"]
    assert "Nenhum objeto em destaque" in unquote(r.headers["X-Mensagem"])


def test_remove_background_rejects_invalid_color(client, fake_model):
    r = remove_bg(client, ("a.jpg", make_photo((300, 200))), background="cor", background_color="vermelho")
    assert r.status_code == 400


# --- Prévia e correção do remover fundo ---

def png_bytes(img):
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def scene():
    """Fundo azul com textura e dois retângulos amarelos; o 'modelo' só achou o da esquerda."""
    size = (400, 300)
    channel = lambda low, seed: noise(size, 12, seed).point(lambda v: low + v // 4)
    img = Image.merge("RGB", [channel(10, 1), channel(80, 2), channel(160, 3)])
    draw = ImageDraw.Draw(img)
    draw.rectangle((40, 60, 160, 240), fill=(240, 200, 30))
    draw.rectangle((240, 60, 360, 240), fill=(235, 195, 40))
    mask = Image.new("L", size, 0)
    ImageDraw.Draw(mask).rectangle((40, 60, 160, 240), fill=255)
    return img, mask


def dot(size, center, radius=8, scale=1):
    """Traço como a prévia manda: RGBA colorido e semitransparente, no tamanho da tela."""
    layer = Image.new("RGBA", (round(size[0] * scale), round(size[1] * scale)))
    x, y, r = center[0] * scale, center[1] * scale, radius * scale
    ImageDraw.Draw(layer).ellipse((x - r, y - r, x + r, y + r), fill=(40, 200, 80, 120))
    return layer


def refine(client, img, mask, strokes, mode):
    return client.post("/background/refine", content_type="multipart/form-data", data={
        "modo": mode, "file": (io.BytesIO(png_bytes(img)), "cena.png"),
        "mascara": (io.BytesIO(png_bytes(mask)), "mascara.png"),
        "traco": (io.BytesIO(png_bytes(strokes)), "traco.png")})


def region(mask, box):
    x0, y0, x1, y1 = box
    return np.asarray(mask)[y0:y1 + 1, x0:x1 + 1]


def test_background_mask_route_returns_the_model_mask(client, fake_model):
    r = client.post("/background/mask", data={"action": "remove-background", "file": (io.BytesIO(make_photo((600, 400))), "f.jpg")},
                    content_type="multipart/form-data")
    mask = opened(r.data)
    assert r.mimetype == "image/png" and mask.mode == "L" and mask.size == (600, 400)
    assert mask.getpixel((300, 200)) == 255 and mask.getpixel((0, 0)) == 0


def background_restored(new):
    """Fração do fundo azul (fora dos dois retângulos) que foi restaurada por engano."""
    outside = np.asarray(new).copy()
    outside[58:243, 38:163] = 0
    outside[58:243, 238:363] = 0
    return (outside > 127).mean()


def test_refine_restores_the_rest_of_a_part_the_model_cut(client):
    img, mask = scene()
    ImageDraw.Draw(mask).rectangle((115, 100, 160, 200), fill=0)  # o modelo cortou um pedaço do retângulo
    # um clique só, numa prévia com metade do tamanho da imagem
    new = opened(refine(client, img, mask, dot(img.size, (138, 150), radius=10, scale=0.5), "restaurar").data)
    assert (region(new, (117, 102, 158, 198)) > 127).mean() > 0.95
    assert background_restored(new) < 0.02
    assert (region(new, (242, 62, 358, 238)) == 0).all()  # o outro retângulo não foi tocado


def test_refine_erases_leftover_background_but_keeps_the_object(client):
    img, mask = scene()
    ImageDraw.Draw(mask).rectangle((170, 100, 225, 200), fill=255)  # pedaço de fundo que ficou
    new = opened(refine(client, img, mask, dot(img.size, (198, 150)), "apagar").data)
    assert (region(new, (172, 102, 223, 198)) < 128).mean() > 0.95
    assert (region(new, (42, 62, 158, 238)) == 255).all()


def test_refine_restores_a_whole_object_the_model_missed_painted_over(client):
    img, _ = scene()
    nothing = Image.new("L", img.size, 0)  # o modelo não achou nenhum dos dois
    stroke = Image.new("RGBA", img.size)
    ImageDraw.Draw(stroke).line([(70, 90), (130, 210)], fill=(40, 200, 80, 255), width=16)
    new = opened(refine(client, img, nothing, stroke, "restaurar").data)
    assert (region(new, (42, 62, 158, 238)) > 127).mean() > 0.95
    assert (region(new, (242, 62, 358, 238)) == 0).all()  # o outro retângulo, que não foi pintado, não muda
    assert background_restored(new) < 0.02


def test_refine_makes_a_half_transparent_part_solid(client):
    img, mask = scene()
    ImageDraw.Draw(mask).rectangle((240, 60, 360, 240), fill=150)  # o modelo ficou em dúvida no da direita
    stroke = Image.new("RGBA", img.size)
    ImageDraw.Draw(stroke).line([(260, 90), (340, 210)], fill=(40, 200, 80, 255), width=16)  # pintado por dentro dele
    new = opened(refine(client, img, mask, stroke, "restaurar").data)
    assert (region(new, (244, 64, 356, 236)) >= 250).mean() > 0.95


def test_refine_erases_a_half_transparent_leftover(client):
    img, mask = scene()
    ImageDraw.Draw(mask).rectangle((170, 100, 225, 200), fill=110)  # fundo que ficou meio visível
    new = opened(refine(client, img, mask, dot(img.size, (198, 150)), "apagar").data)
    assert (region(new, (174, 104, 221, 196)) <= 5).mean() > 0.95
    assert (region(new, (42, 62, 158, 238)) == 255).all()


def test_refine_does_not_follow_a_stroke_that_spills_over(client):
    img, mask = scene()
    ImageDraw.Draw(mask).rectangle((115, 100, 160, 200), fill=0)
    # traço grosso que começa no pedaço que faltou e passa um raio de pincel para o fundo azul
    stroke = Image.new("RGBA", img.size)
    ImageDraw.Draw(stroke).line([(130, 150), (170, 150)], fill=(40, 200, 80, 255), width=16)
    new = np.asarray(opened(refine(client, img, mask, stroke, "restaurar").data))
    assert (new[102:199, 117:159] > 127).mean() > 0.95
    assert (new[143:158, 168:186] > 127).mean() < 0.5  # o fundo por onde o traço vazou não volta


def test_refine_without_smart_detection_changes_exactly_what_was_painted(client):
    img, mask = scene()
    stroke = Image.new("RGBA", img.size)
    # um pedaço do retângulo que o modelo não achou: com a detecção, ele voltaria inteiro
    ImageDraw.Draw(stroke).rectangle((250, 100, 290, 140), fill=(40, 200, 80, 255))
    r = client.post("/background/refine", content_type="multipart/form-data", data={
        "modo": "restaurar", "inteligente": "0", "file": (io.BytesIO(png_bytes(img)), "cena.png"),
        "mascara": (io.BytesIO(png_bytes(mask)), "mascara.png"), "traco": (io.BytesIO(png_bytes(stroke)), "traco.png")})
    new = np.asarray(opened(r.data))
    assert (new[100:141, 250:291] == 255).all()        # o que foi pintado voltou
    changed = new != np.asarray(mask)
    changed[100:141, 250:291] = False
    assert not changed.any()                           # e nada além disso mudou


def textured(size, rgb, seed):
    """Cor com a textura de uma foto: cada canal varia um pouco em volta do valor dado."""
    return Image.merge("RGB", [noise(size, 12, seed + i).point(lambda v, c=c: min(255, max(0, c - 32 + v // 4)))
                               for i, c in enumerate(rgb)])


@pytest.mark.parametrize("mode", ["restaurar", "apagar"])
def test_refine_does_not_spread_to_a_color_found_only_near_the_stroke(client, mode):
    """Uma cor que só aparece perto do traço (no fundo ao restaurar, no objeto ao apagar) não vai
    junto com a correção. A versão que partia de toda a vizinhança do traço marcada como o lado
    pedido estragava aqui uma área quase 7 vezes maior que o erro."""
    size, restore = (400, 300), mode == "restaurar"
    img = textured(size, (60, 110, 170), 1)  # fundo azul
    obj = Image.new("L", size, 0)
    ImageDraw.Draw(obj).rectangle((60, 60, 220, 240), fill=255)
    stain = Image.new("L", size, 0)
    if restore:  # mancha verde no fundo, ao lado do pedaço do objeto que faltou
        ImageDraw.Draw(stain).ellipse((200, 80, 320, 200), fill=255)
        img.paste(textured(size, (60, 170, 80), 4), mask=stain)
        img.paste(textured(size, (230, 190, 50), 7), mask=obj)
    else:  # mancha vermelha no objeto, ao lado do fundo que sobrou
        ImageDraw.Draw(stain).ellipse((130, 80, 250, 200), fill=255)
        img.paste(textured(size, (230, 190, 50), 7), mask=obj)
        img.paste(textured(size, (200, 60, 60), 10), mask=ImageChops.multiply(stain, obj))
    gt = np.asarray(obj) > 0
    yy, xx = np.ogrid[:size[1], :size[0]]
    error = ((xx - 220) ** 2 + (yy - 140) ** 2 <= 30 ** 2) & (gt if restore else ~gt)
    mask = Image.fromarray(((gt ^ error) * 255).astype(np.uint8))
    stroke = Image.new("RGBA", size)
    ImageDraw.Draw(stroke).line([(205, 132), (210, 148)] if restore else [(232, 132), (238, 148)],
                                fill=(40, 200, 80, 255), width=10)
    new = np.asarray(opened(refine(client, img, mask, stroke, mode).data)) > 127
    assert (new[error] == restore).mean() > 0.95
    assert ((new != (np.asarray(mask) > 127)) & ~error).sum() < 0.05 * error.sum()

@pytest.mark.parametrize("change, message", [
    ({"modo": "pintar"}, "restaurar e apagar"),
    ({"mascara": None}, "Falta a máscara"),
    ({"mascara": Image.new("L", (10, 10))}, "tamanho da imagem"),
    ({"traco": Image.new("RGBA", (400, 300))}, "Pinte por cima"),
])
def test_refine_rejects_bad_requests(client, change, message):
    img, mask = scene()
    data = {"modo": "restaurar", "file": (io.BytesIO(png_bytes(img)), "cena.png"),
            "mascara": mask, "traco": dot(img.size, (300, 150))}
    data.update(change)
    for field in ("mascara", "traco"):
        if data[field] is None:
            del data[field]
        elif isinstance(data[field], Image.Image):
            data[field] = (io.BytesIO(png_bytes(data[field])), f"{field}.png")
    r = client.post("/background/refine", data=data, content_type="multipart/form-data")
    assert r.status_code == 400
    assert message in r.get_data(as_text=True)


def test_download_uses_the_mask_edited_in_the_preview(client, fake_model):
    img, mask = scene()
    r = client.post("/convert", content_type="multipart/form-data", data={
        "action": "remove-background", "file": (io.BytesIO(png_bytes(img)), "cena.png"),
        "mascara": (io.BytesIO(png_bytes(mask)), "mascara.png")})
    alpha = np.asarray(opened(r.data))[..., 3]
    assert np.array_equal(alpha, np.asarray(mask))  # e não o disco do modelo falso


def test_download_needs_one_mask_per_image(client, fake_model):
    img, mask = scene()
    r = client.post("/convert", content_type="multipart/form-data", data={
        "action": "remove-background",
        "file": [(io.BytesIO(png_bytes(img)), "a.png"), (io.BytesIO(png_bytes(img)), "b.png")],
        "mascara": (io.BytesIO(png_bytes(mask)), "mascara.png")})
    assert r.status_code == 400


@pytest.fixture
def model_source(tmp_path, monkeypatch):
    """Um 'modelo' servido de um arquivo local, com o SHA-256 dele."""
    source = tmp_path / "modelo.onnx"
    source.write_bytes(b"modelo de teste" * 1000)
    monkeypatch.setattr(euamopdf, "BG_MODEL_URL", source.as_uri())
    monkeypatch.setattr(euamopdf, "BG_MODEL_SHA256", hashlib.sha256(source.read_bytes()).hexdigest())
    return source


def test_model_download_puts_the_file_in_place(tmp_path, model_source):
    target = tmp_path / "cache" / "modelo.onnx"
    euamopdf.download_bg_model(target)
    assert target.read_bytes() == model_source.read_bytes()
    assert list(target.parent.iterdir()) == [target]  # nada de .part sobrando


def test_model_download_rejects_a_corrupted_file(tmp_path, model_source, monkeypatch):
    monkeypatch.setattr(euamopdf, "BG_MODEL_SHA256", "0" * 64)
    target = tmp_path / "cache" / "modelo.onnx"
    with pytest.raises(euamopdf.UserError, match="corrompido"):
        euamopdf.download_bg_model(target)
    assert list(target.parent.iterdir()) == []


def test_model_download_without_internet(tmp_path, monkeypatch):
    monkeypatch.setattr(euamopdf, "BG_MODEL_URL", (tmp_path / "nao-existe.onnx").as_uri())
    with pytest.raises(euamopdf.UserError, match="internet"):
        euamopdf.download_bg_model(tmp_path / "cache" / "modelo.onnx")


def test_model_goes_to_the_system_cache(monkeypatch, tmp_path):
    monkeypatch.delenv("LOCALAPPDATA", raising=False)
    monkeypatch.setenv("XDG_CACHE_HOME", str(tmp_path))
    assert euamopdf.bg_model_path() == tmp_path / "euamopdf" / "isnet-general-use.onnx"


def test_first_use_warns_about_the_model_download(client, monkeypatch, tmp_path):
    monkeypatch.setattr(euamopdf, "bg_model_path", lambda: tmp_path / "nao-baixado.onnx")
    assert "baixa o modelo de IA" in client.get("/").get_data(as_text=True)
    monkeypatch.setattr(euamopdf, "bg_model_path", lambda: tmp_path)  # "já baixado"
    assert "baixa o modelo de IA" not in client.get("/").get_data(as_text=True)


@pytest.mark.skipif(not euamopdf.bg_model_path().exists(), reason="modelo de remover fundo não baixado")
def test_remove_background_with_the_real_model(client, monkeypatch):
    monkeypatch.setattr(euamopdf, "_bg_session", None)
    scene = Image.open(io.BytesIO(make_photo((1200, 800)))).filter(ImageFilter.GaussianBlur(25))
    subject = Image.new("L", scene.size, 0)
    ImageDraw.Draw(subject).ellipse((400, 200, 800, 600), fill=255)
    scene.paste(Image.new("RGB", scene.size, (230, 60, 40)), mask=subject)  # disco vermelho em destaque
    buf = io.BytesIO()
    scene.save(buf, "PNG")
    alpha = np.asarray(opened(remove_bg(client, ("cena.png", buf.getvalue())).data))[..., 3] > 127
    truth = np.asarray(subject) > 127
    assert (alpha & truth).sum() / (alpha | truth).sum() > 0.9


# --- Metadados ---

XMP_ANTIGO = """<?xpacket begin="\ufeff" id="W5M0MpCehiHzreSzNTczkc9d"?>
<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">
<rdf:Description rdf:about="" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:pdf="http://ns.adobe.com/pdf/1.3/"
    xmlns:xmp="http://ns.adobe.com/xap/1.0/" xmlns:pdfaid="http://www.aiim.org/pdfa/ns/id/"
    pdf:Producer="Produtor antigo" xmp:CreatorTool="Writer" pdfaid:part="2" pdfaid:conformance="B">
  <dc:title><rdf:Alt><rdf:li xml:lang="x-default">Título antigo</rdf:li></rdf:Alt></dc:title>
  <dc:creator><rdf:Seq><rdf:li>Autor antigo</rdf:li></rdf:Seq></dc:creator>
</rdf:Description></rdf:RDF></x:xmpmeta>
<?xpacket end="w"?>"""


def pdf_with_metadata(xmp=None):
    doc = pymupdf.open()
    doc.new_page()
    doc.set_metadata({"title": "Título antigo", "author": "Autor antigo", "producer": "Produtor antigo",
                      "creator": "Writer", "creationDate": "D:20260102030405"})
    if xmp:
        doc.set_xml_metadata(xmp)
    return doc.tobytes()


def edit(client, pdf, **form):
    return post(client, "edit-metadata", ("relatorio.pdf", pdf), **form)


NOVOS = {"meta_title": "Relatório de ação", "meta_author": "João da Silva", "meta_subject": "Contas",
         "meta_keywords": "finanças, 2026", "meta_creator": "", "meta_producer": "EuAmoPDF",
         "meta_created": "2026-01-02T03:04:05", "meta_modified": "2026-09-29T14:30:00"}


def test_metadata_route_shows_current_values(client):
    r = client.post("/metadata", data={"file": (io.BytesIO(pdf_with_metadata()), "a.pdf")}, content_type="multipart/form-data")
    values = r.get_json()
    assert values["meta_title"] == "Título antigo" and values["meta_producer"] == "Produtor antigo"
    assert values["meta_created"] == "2026-01-02T03:04:05" and values["meta_modified"] == ""


def test_edit_metadata_updates_the_info_dictionary(client):
    r = edit(client, pdf_with_metadata(), **NOVOS)
    assert r.headers["Content-Disposition"].endswith("relatorio.pdf")
    meta = pymupdf.open(stream=r.data, filetype="pdf").metadata
    assert (meta["title"], meta["author"], meta["subject"], meta["keywords"], meta["producer"]) == \
        ("Relatório de ação", "João da Silva", "Contas", "finanças, 2026", "EuAmoPDF")
    assert meta["creator"] == ""  # campo esvaziado some
    assert euamopdf.parse_pdf_date(meta["modDate"]) == datetime.fromisoformat("2026-09-29T14:30:00").astimezone()


def test_edit_metadata_keeps_xmp_in_sync_and_the_rest_of_it(client):
    r = edit(client, pdf_with_metadata(XMP_ANTIGO), **NOVOS)
    xmp = pymupdf.open(stream=r.data, filetype="pdf").get_xml_metadata()
    assert "Relatório de ação" in xmp and "João da Silva" in xmp and "EuAmoPDF" in xmp
    assert "antigo" not in xmp and "Writer" not in xmp  # nada do valor velho, nem como atributo
    assert xmp.count("Relatório de ação") == 1
    assert 'pdfaid:part="2"' in xmp  # o que não foi editado continua
    assert "2026-09-29T14:30:00" in xmp


@pytest.mark.parametrize("raw", ["D:20260102030405", "D:20260102030400Z", "D:20260102030405+05'30'",
                                 "20260102030405", "2026-01-02T03:04:05", "Tue Jan 02 03:04:05 2026"])  # sem o D: etc.
def test_pdf_date_saved_unchanged_stays_as_it_was(client, raw):
    doc = pymupdf.open()
    doc.new_page()
    doc.set_metadata({"creationDate": raw})
    pdf = doc.tobytes()
    form = {k: v for k, v in client.post("/metadata", data={"file": (io.BytesIO(pdf), "a.pdf")},
                                         content_type="multipart/form-data").get_json().items() if k.startswith("meta_")}
    form["meta_created"] = form["meta_created"].removesuffix(":00")  # o navegador omite os segundos zerados
    r = edit(client, pdf, meta_title="Outro título", **{k: v for k, v in form.items() if k != "meta_title"})
    meta = pymupdf.open(stream=r.data, filetype="pdf").metadata
    assert meta["title"] == "Outro título" and meta["creationDate"] == raw  # sem ganhar nem trocar o fuso


def test_pdf_metadata_only_in_xmp_is_shown_and_kept(client):
    doc = pymupdf.open()
    doc.new_page()
    doc.set_xml_metadata(XMP_ANTIGO.replace('pdfaid:part', 'xmp:CreateDate="2026-01-02T03:04:05Z" pdfaid:part'))
    pdf = doc.tobytes()
    form = {k: v for k, v in client.post("/metadata", data={"file": (io.BytesIO(pdf), "a.pdf")},
                                         content_type="multipart/form-data").get_json().items() if k.startswith("meta_")}
    assert form["meta_title"] == "Título antigo" and form["meta_producer"] == "Produtor antigo"
    assert form["meta_created"] == datetime(2026, 1, 2, 3, 4, 5, tzinfo=timezone.utc).astimezone().strftime("%Y-%m-%dT%H:%M:%S")
    r = edit(client, pdf, **form | {"meta_subject": "Contas"})
    out = pymupdf.open(stream=r.data, filetype="pdf")
    assert out.metadata["title"] == "Título antigo" and "Título antigo" in out.get_xml_metadata()
    assert "2026-01-02T03:04:05Z" in out.get_xml_metadata()  # a data não mexida fica como estava


def test_edit_metadata_does_not_create_xmp(client):
    r = edit(client, pdf_with_metadata(), **NOVOS)
    assert pymupdf.open(stream=r.data, filetype="pdf").get_xml_metadata() == ""


def test_remove_all_metadata(client):
    r = edit(client, pdf_with_metadata(XMP_ANTIGO), meta_strip="1")
    doc = pymupdf.open(stream=r.data, filetype="pdf")
    assert not any(v for k, v in doc.metadata.items() if k != "format")
    assert doc.get_xml_metadata() == ""
    assert "removidos" in unquote(r.headers["X-Mensagem"])


XMP_FOTO = ('<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">'
            '<rdf:Description rdf:about="" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:exif="http://ns.adobe.com/exif/1.0/" '
            'exif:GPSLatitude="23,33.0S" exif:GPSLongitude="46,38.0W"><dc:title><rdf:Alt>'
            '<rdf:li xml:lang="en">Old photo</rdf:li><rdf:li xml:lang="x-default">Foto antiga</rdf:li></rdf:Alt></dc:title>'
            '</rdf:Description></rdf:RDF></x:xmpmeta>')


def phone_photo(fmt="JPEG"):
    """Foto como a de um celular: câmera, orientação, localização GPS e XMP."""
    exif = Image.Exif()
    exif[0x010F], exif[0x0110], exif[0x0112] = "Samsung", "Galaxy", 6
    exif[0x8825] = {1: "S", 2: (23.0, 33.0, 0.0), 3: "W", 4: (46.0, 38.0, 0.0)}
    img = opened(make_photo((300, 200)))
    buf = io.BytesIO()
    if fmt == "PNG":
        from PIL import PngImagePlugin
        info = PngImagePlugin.PngInfo()
        info.add_itxt("XML:com.adobe.xmp", XMP_FOTO)
        img.save(buf, "PNG", exif=exif, pnginfo=info)
    elif fmt == "JPEG":  # segmento XMP montado à mão: o parâmetro xmp= do JPEG só existe no Pillow 11
        img.save(buf, "JPEG", exif=exif, quality=90)
        data, xmp = buf.getvalue(), b"http://ns.adobe.com/xap/1.0/\x00" + XMP_FOTO.encode()
        return data[:2] + b"\xff\xe1" + (len(xmp) + 2).to_bytes(2, "big") + xmp + data[2:]
    else:
        img.save(buf, fmt, exif=exif, xmp=XMP_FOTO.encode(), quality=90)
    return buf.getvalue()


def image_metadata(client, data, name):
    return client.post("/metadata", data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data").get_json()


def same_pixels(a, b):
    return np.array_equal(np.asarray(opened(a).convert("RGBA")), np.asarray(opened(b).convert("RGBA")))


FOTO = {"img_title": "Pôr do sol", "img_description": "Na praia", "img_author": "João", "img_copyright": "© João 2026",
        "img_keywords": "praia, férias", "img_software": "", "img_taken": "2026-01-15T18:20:00"}


@pytest.mark.parametrize("fmt, ext", [("JPEG", "jpg"), ("PNG", "png"), ("WEBP", "webp")])
def test_image_metadata_is_read_and_edited_without_touching_the_image(client, fmt, ext):
    src = phone_photo(fmt)
    before = image_metadata(client, src, f"f.{ext}")
    assert before["tipo"] == "imagem" and before["camera"] == "Samsung Galaxy"
    assert before["img_title"] == "Foto antiga" and before["localizacao"] == "-23.55000, -46.63333"
    r = post(client, "edit-metadata", (f"foto.{ext}", src), **FOTO)
    assert r.headers["Content-Disposition"].endswith(f"foto.{ext}")
    assert same_pixels(src, r.data)
    after = image_metadata(client, r.data, f"f.{ext}")
    assert {k: after[k] for k in FOTO if FOTO[k]} == {k: v for k, v in FOTO.items() if v}
    assert after["localizacao"] == "-23.55000, -46.63333"  # sem pedir, a localização fica
    assert opened(r.data).getexif()[0x0112] == 6


@pytest.mark.parametrize("fmt, ext", [("JPEG", "jpg"), ("PNG", "png"), ("WEBP", "webp")])
def test_image_date_taken_goes_to_exif_and_leaves_no_empty_block(client, fmt, ext):
    src = make_photo((200, 150), fmt=fmt)  # sem EXIF nem XMP: a data só pode ir para o bloco Exif novo
    r = post(client, "edit-metadata", (f"f.{ext}", src), img_taken="2026-05-06T07:08:09")
    assert opened(r.data).getexif().get_ifd(0x8769)[0x9003] == "2026:05:06 07:08:09"
    r = post(client, "edit-metadata", (f"f.{ext}", r.data), img_taken="")
    assert 0x8769 not in opened(r.data).getexif()


@pytest.mark.parametrize("creation", ["ter 29 set 2026 14:21:03", "Tue 29 Sep 2026 02:21:03 PM BRT",
                                      "Tue, 29 Sep 2026 14:21:03 GMT", "Tue Sep 29 14:21:03 2026"])
def test_png_creation_time_in_words_is_read_and_kept(client, creation):
    from PIL import PngImagePlugin
    info = PngImagePlugin.PngInfo()  # como o gnome-screenshot grava (idioma do sistema) e como a norma do PNG pede
    info.add_text("Software", "gnome-screenshot")
    info.add_text("Creation Time", creation)
    buf = io.BytesIO()
    opened(make_photo((50, 40), fmt="PNG")).save(buf, "PNG", pnginfo=info)
    form = {k: v for k, v in image_metadata(client, buf.getvalue(), "f.png").items() if k.startswith("img_")}
    assert form["img_taken"] == "2026-09-29T14:21:03" and form["img_software"] == "gnome-screenshot"
    r = post(client, "edit-metadata", ("f.png", buf.getvalue()), **form | {"img_title": "Tela"})
    assert image_metadata(client, r.data, "f.png")["img_taken"] == "2026-09-29T14:21:03"  # editar não apaga a data


@pytest.mark.parametrize("fmt, ext", [("JPEG", "jpg"), ("PNG", "png"), ("WEBP", "webp")])
def test_image_metadata_saved_unchanged_adds_no_exif_block(client, fmt, ext):
    src = phone_photo(fmt)  # EXIF com câmera e GPS, sem o bloco Exif (0x8769)
    form = {k: v for k, v in image_metadata(client, src, f"f.{ext}").items() if k.startswith("img_")}
    r = post(client, "edit-metadata", (f"f.{ext}", src), **form)
    assert 0x8769 not in opened(r.data).getexif()


@pytest.mark.parametrize("fmt, ext", [("JPEG", "jpg"), ("PNG", "png"), ("WEBP", "webp")])
def test_image_location_is_removed_from_exif_and_xmp(client, fmt, ext):
    r = post(client, "edit-metadata", (f"foto.{ext}", phone_photo(fmt)), img_strip_gps="1", **FOTO)
    img = opened(r.data)
    assert not img.getexif().get_ifd(0x8825)
    xmp = img.info.get("xmp") or img.info.get("XML:com.adobe.xmp") or b""
    assert b"GPS" not in (xmp if isinstance(xmp, bytes) else xmp.encode())
    assert image_metadata(client, r.data, f"f.{ext}")["localizacao"] == ""


def photo_with_something_after_it():
    """Foto com outra imagem grudada depois do fim, com a localização dela: é o que guardam as fotos em
    movimento (um vídeo), o MPO de alguns celulares e o trailer das fotos da Samsung."""
    extra = make_photo((40, 30), exif=exif_with_location())
    return make_photo((200, 150)) + extra, extra


@pytest.mark.parametrize("form", [{"img_strip_gps": "1"}, {"meta_strip": "1"}])
def test_removing_the_location_drops_what_comes_after_the_jpeg(client, form):
    photo, extra = photo_with_something_after_it()
    r = post(client, "edit-metadata", ("foto.jpg", photo), **form)
    assert extra not in r.data
    assert same_pixels(r.data, photo)
    assert "depois da foto" in unquote(r.headers["X-Mensagem"])


@pytest.mark.parametrize("options", [{}, {"progressive": True}, {"restart_marker_rows": 1}])
def test_the_end_of_the_jpeg_is_found_exactly(options):
    """O JPEG progressivo tem várias varreduras e segmentos no meio, e os marcadores RST ficam dentro
    dos dados: cortar antes do fim estragaria a foto."""
    buf = io.BytesIO()
    opened(make_photo((400, 300))).save(buf, "JPEG", quality=90, **options)
    jpeg = buf.getvalue()
    assert euamopdf.jpeg_image_end(jpeg + b"depois", jpeg.index(b"\xff\xda")) == len(jpeg)


def test_editing_the_title_keeps_what_comes_after_the_jpeg(client):
    photo, extra = photo_with_something_after_it()
    r = post(client, "edit-metadata", ("foto.jpg", photo), img_title="Praia")
    assert r.data.endswith(extra)


@pytest.mark.parametrize("fmt, ext", [("JPEG", "jpg"), ("PNG", "png"), ("WEBP", "webp")])
def test_remove_all_image_metadata_keeps_orientation(client, fmt, ext):
    src = phone_photo(fmt)
    r = post(client, "edit-metadata", (f"foto.{ext}", src), meta_strip="1")
    img = opened(r.data)
    assert dict(img.getexif()) == {0x0112: 6}  # sem ela, a foto apareceria deitada
    assert not (img.info.get("xmp") or img.info.get("XML:com.adobe.xmp"))
    assert same_pixels(src, r.data)


def iptc_photo():
    """JPEG com IPTC como o do Photoshop: resumo MD5, outro recurso no bloco e textos em Latin-1."""
    def dataset(number, value, record=2):
        return bytes((0x1C, record, number)) + len(value).to_bytes(2, "big") + value
    iptc = (dataset(0, b"\x00\x04") + dataset(5, "Título IPTC".encode("latin-1")) + dataset(25, b"praia") + dataset(25, b"sol")
            + dataset(55, b"20260102") + dataset(60, b"030405-0300") + dataset(80, b"Ana")
            + dataset(90, "São Paulo".encode("latin-1")) + dataset(120, b"Legenda"))
    def resource(number, data):
        return b"8BIM" + number.to_bytes(2, "big") + b"\x00\x00" + len(data).to_bytes(4, "big") + data + b"\x00" * (len(data) & 1)
    body = b"Photoshop 3.0\x00" + resource(0x03ED, bytes(16)) + resource(0x0404, iptc) + resource(0x0425, hashlib.md5(iptc).digest())
    jpeg = make_photo((60, 40), fmt="JPEG")
    return jpeg[:2] + b"\xff\xed" + (len(body) + 2).to_bytes(2, "big") + body + jpeg[2:]


def test_jpeg_iptc_is_read_and_updated_keeping_the_rest(client):
    from PIL import IptcImagePlugin
    src = iptc_photo()
    form = {k: v for k, v in image_metadata(client, src, "f.jpg").items() if k.startswith("img_")}
    assert (form["img_title"], form["img_description"], form["img_author"], form["img_keywords"], form["img_taken"]) == \
        ("Título IPTC", "Legenda", "Ana", "praia, sol", "2026-01-02T03:04:05")
    r = post(client, "edit-metadata", ("f.jpg", src), **form | {"img_title": "Pôr do sol", "img_description": ""})
    img = opened(r.data)
    iptc = IptcImagePlugin.getiptcinfo(img)
    assert iptc[(2, 5)] == "Pôr do sol".encode() and (2, 120) not in iptc  # senão a legenda voltaria na próxima leitura
    assert iptc[(2, 25)] == [b"praia", b"sol"] and iptc[(2, 55)] == b"20260102"
    assert iptc[(2, 90)] == "São Paulo".encode() and iptc[(1, 90)] == b"\x1b%G"  # o que não é do formulário fica, em UTF-8
    body = next(data for marker, data in img.applist if marker == "APP13")
    resources = euamopdf.photoshop_resources(body[len(b"Photoshop 3.0\x00"):])
    assert [r[0] for r in resources] == [b"\x03\xed", b"\x04\x04", b"\x04\x25"]
    assert resources[2][2] == hashlib.md5(resources[1][2]).digest()
    assert image_metadata(client, r.data, "f.jpg")["img_description"] == ""
    assert same_pixels(src, r.data)
    assert "photoshop" not in opened(post(client, "edit-metadata", ("f.jpg", phone_photo()), **FOTO).data).info  # sem IPTC, não ganha um


def test_jpeg_image_data_stays_byte_for_byte_identical(client):
    src = phone_photo("JPEG")
    out = post(client, "edit-metadata", ("foto.jpg", src), **FOTO).data
    assert out[out.index(b"\xff\xda"):] == src[src.index(b"\xff\xda"):]  # nada recomprimido


@pytest.mark.parametrize("lossless, mode", [(False, "RGB"), (True, "RGBA")])  # com perda e transparência já vem estendido
def test_simple_webp_gets_the_extended_header_for_metadata(client, lossless, mode):
    buf = io.BytesIO()
    opened(make_photo((120, 80), fmt="PNG", mode=mode)).save(buf, "WEBP", lossless=lossless)
    src = buf.getvalue()
    assert src[12:16] in (b"VP8 ", b"VP8L")  # sem VP8X
    out = post(client, "edit-metadata", ("icone.webp", src), img_author="João").data
    assert out[12:16] == b"VP8X"
    assert image_metadata(client, out, "i.webp")["img_author"] == "João"
    assert same_pixels(src, out)


def test_edit_metadata_rejects_invalid_date(client):
    assert edit(client, pdf_with_metadata(), meta_created="ontem").status_code == 400


# --- Converter de PDF ---

def test_pdf_to_jpg_converts_every_page(client):
    r = post(client, "pdf-to-jpg", ("doc.pdf", make_pdf(3)), dpi="72")
    files = unzip(r.data)
    assert sorted(files) == ["doc_pag_1.jpg", "doc_pag_2.jpg", "doc_pag_3.jpg"]
    assert Image.open(io.BytesIO(files["doc_pag_1.jpg"])).size == (595, 842)


def test_pdf_to_jpg_single_page_returns_jpg(client):
    r = post(client, "pdf-to-jpg", ("doc.pdf", make_pdf(1)))
    assert Image.open(io.BytesIO(r.data)).format == "JPEG"


def test_pdf_to_word(client):
    r = post(client, "pdf-to-word", ("doc.pdf", make_pdf(2)))
    assert r.status_code == 200
    assert b"Pagina 2" in unzip(r.data)["word/document.xml"]


def test_pdf_to_word_keeps_the_table_column_widths(client):
    """O pdf2docx grava a largura certa em cada célula, mas a grade da tabela com as colunas iguais, e
    os editores desenham pela grade: a coluna estreitada quebrava o texto, e a altura exata o cortava."""
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn
    doc = pymupdf.open()
    page = doc.new_page()
    for r in range(3):
        x = 72
        for c, width in enumerate((40, 260, 90)):
            cell = pymupdf.Rect(x, 72 + r * 24, x + width, 96 + r * 24)
            page.draw_rect(cell, color=(0, 0, 0), width=0.5)
            page.insert_text(cell.bl + (4, -7), f"L{r}C{c}")
            x += width
    r = post(client, "pdf-to-word", ("tabela.pdf", doc.tobytes()))
    table = docx.Document(io.BytesIO(r.data)).tables[0]._tbl
    grid = [col.get(qn("w:w")) for col in table.tblGrid.findall(qn("w:gridCol"))]
    cells = [tc.tcPr.find(qn("w:tcW")).get(qn("w:w")) for tc in table.findall(qn("w:tr"))[0].findall(qn("w:tc"))]
    assert len(set(cells)) == 3  # as três larguras do PDF
    assert grid == cells


@pytest.fixture
def fake_pdf2docx(monkeypatch):
    """pdf2docx simulado: a conversão n grava um .docx com o texto texts[n]. Devolve (opções de cada
    conversão, texts)."""
    docx = pytest.importorskip("docx")
    import pdf2docx
    calls, texts = [], []

    class Converter:
        def __init__(self, path):
            pass

        def convert(self, out, **options):
            calls.append(options)
            document = docx.Document()
            document.add_paragraph(texts[len(calls) - 1])
            document.save(out)

        def close(self):
            pass

    monkeypatch.setattr(pdf2docx, "Converter", Converter)
    return calls, texts


def test_pdf_to_word_tries_again_when_text_is_lost(client, fake_pdf2docx):
    """A detecção de tabelas sem bordas do pdf2docx descartava o texto de um PDF do "Microsoft Print to
    PDF" (41% das palavras): sem ela, a segunda tentativa traz o texto de volta."""
    calls, texts = fake_pdf2docx
    texts += ["Pagina 1", "Pagina 1 Pagina 2 Pagina 3"]
    r = post(client, "pdf-to-word", ("doc.pdf", make_pdf(3)))
    assert calls == [{}, {"parse_stream_table": False}]
    assert b"Pagina 3" in unzip(r.data)["word/document.xml"]
    assert "X-Mensagem" not in r.headers


def test_pdf_to_word_warns_when_text_is_still_missing(client, fake_pdf2docx):
    calls, texts = fake_pdf2docx
    texts += ["Pagina 1", "Pagina 1 Pagina 2"]
    r = post(client, "pdf-to-word", ("doc.pdf", make_pdf(3)))
    assert b"Pagina 2" in unzip(r.data)["word/document.xml"]  # fica a tentativa com mais texto
    assert "Parte do texto" in unquote(r.headers["X-Mensagem"])


def test_pdf_to_word_uses_the_real_font_name(client):
    """O "Microsoft Print to PDF" chama todas as fontes de CIDFont+F1, F2...; o editor não conhece esse
    nome e usava outra fonte, mais larga, que cortava o texto. O nome real está no arquivo da fonte."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="F1", fontbuffer=pymupdf.Font("tiro").buffer)  # Nimbus Roman, embutida
    page.insert_text((72, 100), "Texto com a fonte renomeada", fontname="F1", fontsize=12)
    pdf = pymupdf.open(stream=doc.tobytes())
    for xref in range(1, pdf.xref_length()):
        for key in ("BaseFont", "FontName"):
            if pdf.xref_get_key(xref, key)[0] == "name":
                pdf.xref_set_key(xref, key, "/CIDFont+F1")
    xml = unzip(post(client, "pdf-to-word", ("doc.pdf", pdf.tobytes())).data)["word/document.xml"].decode()
    assert set(re.findall(r'w:ascii="([^"]+)"', xml)) == {"Nimbus Roman"}


def test_pdf_to_word_keeps_what_is_in_annotations_and_fields(client):
    """O pdf2docx só tira as imagens do conteúdo da página: o selo visível de uma assinatura (gov.br),
    que fica num campo ou numa anotação, sumia. E o checkbox marcado virava "3", o código do ✓ na
    fonte ZapfDingbats."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Aceito os termos")
    box = pymupdf.Widget()
    box.field_type, box.field_name, box.field_value = pymupdf.PDF_WIDGET_TYPE_CHECKBOX, "aceito", True
    box.rect = pymupdf.Rect(190, 88, 204, 102)
    page.add_widget(box)
    page.add_stamp_annot(pymupdf.Rect(72, 200, 252, 290), stamp=make_image())
    xml = unzip(post(client, "pdf-to-word", ("doc.pdf", doc.tobytes())).data)["word/document.xml"].decode()
    assert xml.count("<pic:pic") == 1
    assert "✓" in xml and not re.search(r"<w:t[^>]*>\s*3\s*</w:t>", xml)


def test_pdf_to_word_converts_once_when_no_text_is_lost(client, monkeypatch):
    from pdf2docx import Converter
    calls = []
    convert = Converter.convert
    monkeypatch.setattr(Converter, "convert", lambda self, *args, **kwargs: calls.append(kwargs) or convert(self, *args, **kwargs))
    assert post(client, "pdf-to-word", ("doc.pdf", make_pdf(3))).status_code == 200
    assert calls == [{}]


def make_table_pdf(rows):
    """PDF com uma tabela de bordas desenhadas, como as geradas por planilhas."""
    doc = pymupdf.open()
    page = doc.new_page()
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = pymupdf.Rect(72 + c * 120, 72 + r * 24, 192 + c * 120, 96 + r * 24)
            page.draw_rect(cell, color=(0, 0, 0), width=0.5)
            page.insert_text(cell.bl + (4, -7), value)
    return doc.tobytes()


def test_pdf_to_excel_extracts_tables(client):
    openpyxl = pytest.importorskip("openpyxl")
    rows = [["Produto", "Preço"], ["Café", "12,50"], ["Pão", "0,75"]]
    r = post(client, "pdf-to-excel", ("tabela.pdf", make_table_pdf(rows)))
    assert r.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(r.data)).worksheets[0]
    assert [[cell.value for cell in row] for row in sheet.iter_rows()] == [["Produto", "Preço"], ["Café", 12.5], ["Pão", 0.75]]


def draw_table(page, rows, top):
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = pymupdf.Rect(72 + c * 150, top + r * 22, 222 + c * 150, top + 22 + r * 22)
            page.draw_rect(cell, color=(0, 0, 0), width=0.5)
            page.insert_text(cell.bl + (4, -7), value, fontsize=9)


def test_pdf_to_excel_joins_a_table_that_continues_on_the_next_page(client):
    """Uma tabela longa, que passa para a página seguinte repetindo o cabeçalho, virava uma aba por
    página."""
    openpyxl = pytest.importorskip("openpyxl")
    doc = pymupdf.open()
    for p in range(2):
        draw_table(doc.new_page(), [["Item", "Quantidade"]] + [[f"Item {p * 30 + i + 1}", "1"] for i in range(30)], 60)
    wb = openpyxl.load_workbook(io.BytesIO(post(client, "pdf-to-excel", ("longa.pdf", doc.tobytes())).data))
    assert len(wb.worksheets) == 1
    values = [row for row in wb.worksheets[0].iter_rows(values_only=True)]
    assert values[0] == ("Item", "Quantidade") and len(values) == 61 and values[-1] == ("Item 60", "1")


def test_pdf_to_excel_keeps_different_tables_apart(client):
    openpyxl = pytest.importorskip("openpyxl")
    doc = pymupdf.open()
    page = doc.new_page()
    draw_table(page, [["Produto", "Preço"], ["Café", "12,50"]], 60)
    draw_table(page, [["Cidade", "Estado"], ["Recife", "PE"]], 400)
    draw_table(doc.new_page(), [["Nome", "Idade"], ["Ana", "30"]], 600)  # nem no alto da página seguinte
    wb = openpyxl.load_workbook(io.BytesIO(post(client, "pdf-to-excel", ("tabelas.pdf", doc.tobytes())).data))
    assert len(wb.worksheets) == 3


def test_pdf_to_excel_finds_tables_without_borders(client):
    """Extratos de banco alinham as colunas sem desenhar linhas: a busca por linhas não achava nada, e
    a mensagem mandava passar OCR num PDF que já tem texto."""
    openpyxl = pytest.importorskip("openpyxl")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((60, 60), "Extrato da conta corrente - outubro de 2026", fontsize=12)
    rows = [("Data", "Descrição", "Valor", "Saldo"), ("01/10/2026", "PIX recebido", "1.234,56", "5.000,00"),
            ("02/10/2026", "Boleto pago", "-250,00", "4.750,00"), ("03/10/2026", "Tarifa", "-12,90", "4.737,10")]
    for i, row in enumerate(rows):
        for x, text in zip((60, 150, 330, 430), row):
            page.insert_text((x, 100 + 18 * i), text, fontsize=10)
    r = post(client, "pdf-to-excel", ("extrato.pdf", doc.tobytes()))
    assert r.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(r.data)).worksheets[0]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert [datetime(2026, 10, 1), "PIX recebido", 1234.56, 5000.0] in values
    assert not any(all(v is None for v in row) for row in values)  # sem as linhas vazias da detecção
    assert "linhas" in unquote(r.headers["X-Mensagem"])  # avisa que a tabela pode precisar de ajuste


def test_pdf_to_excel_writes_amounts_and_dates_as_such(client):
    """Valores e datas entravam como texto, e o Excel não somava nem ordenava por data. Inteiros
    continuam texto: contas, CEPs e documentos têm zeros à esquerda."""
    openpyxl = pytest.importorskip("openpyxl")
    rows = [["Data", "Valor", "CPF", "Conta"],
            ["01/10/2026", "1.234,56", "123.456.789-00", "000123"],
            ["31/12/2026", "-12,90", "01001-000", "R$ 10,00"]]
    r = post(client, "pdf-to-excel", ("extrato.pdf", make_table_pdf(rows)))
    sheet = openpyxl.load_workbook(io.BytesIO(r.data)).worksheets[0]
    assert [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)] == [
        [datetime(2026, 10, 1), 1234.56, "123.456.789-00", "000123"],
        [datetime(2026, 12, 31), -12.9, "01001-000", 10.0]]
    assert sheet["B2"].number_format == "#,##0.00" and sheet["A2"].number_format == "dd/mm/yyyy"


def test_pdf_to_excel_without_tables(client):
    r = post(client, "pdf-to-excel", ("texto.pdf", make_pdf()))
    assert r.status_code == 400
    assert "Nenhuma tabela" in r.get_data(as_text=True)


def test_pdf_to_ppt_one_slide_per_page(client):
    pptx = pytest.importorskip("pptx")
    doc = pymupdf.open()
    doc.new_page(width=842, height=595)  # A4 deitada
    doc.new_page(width=595, height=842)  # A4 em pé: deve caber sem distorcer
    r = post(client, "pdf-to-ppt", ("slides.pdf", doc.tobytes()))
    assert r.status_code == 200
    prs = pptx.Presentation(io.BytesIO(r.data))
    assert len(prs.slides) == 2
    assert prs.slide_width > prs.slide_height
    picture = prs.slides[1].shapes[0]
    assert picture.height == prs.slide_height
    assert abs(picture.width / picture.height - 595 / 842) < 0.01


# --- Converter para PDF ---

def test_images_to_pdf_one_a4_page_per_image(client):
    r = post(client, "jpg-to-pdf", ("a.png", make_image()), ("b.jpg", make_image(fmt="JPEG")))
    doc = pymupdf.open(stream=r.data, filetype="pdf")
    assert doc.page_count == 2
    assert [round(x) for x in doc[0].rect] == [0, 0, 842, 595]  # imagem deitada, A4 deitada


def test_images_to_pdf_transparency_becomes_white(client):
    r = post(client, "jpg-to-pdf", ("t.png", make_image("RGBA", color=(0, 0, 0, 0))))
    page = pymupdf.open(stream=r.data, filetype="pdf")[0]
    assert page.get_pixmap(dpi=10).pixel(10, 10) == (255, 255, 255)


def test_images_to_pdf_follows_exif_orientation(client):
    img = Image.new("RGB", (400, 300))
    exif = img.getexif()
    exif[0x0112] = 6  # foto de celular: gravada deitada, deve ser exibida em pé
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif)
    r = post(client, "jpg-to-pdf", ("foto.jpg", buf.getvalue()))
    page = pymupdf.open(stream=r.data, filetype="pdf")[0]
    assert page.rect.height > page.rect.width


def test_images_to_pdf_keeps_every_page_of_a_tiff(client):
    """TIFF de scanner ou fax costuma ter várias páginas; só a primeira ia para o PDF."""
    pages = [Image.new("L", (200, 300), shade) for shade in (40, 128, 220)]
    buf = io.BytesIO()
    pages[0].save(buf, "TIFF", save_all=True, append_images=pages[1:])
    r = post(client, "jpg-to-pdf", ("digitalizado.tif", buf.getvalue()))
    assert pymupdf.open(stream=r.data).page_count == 3


def test_images_to_pdf_keeps_16_bit_images(client):
    """PNG de 16 bits em tons de cinza saía como uma página branca."""
    gradient = np.tile(np.linspace(0, 65535, 600), (400, 1)).astype(np.uint16)
    buf = io.BytesIO()
    Image.fromarray(gradient).save(buf, "PNG")
    r = post(client, "jpg-to-pdf", ("cinza.png", buf.getvalue()))
    pix = pymupdf.open(stream=r.data)[0].get_pixmap(dpi=30)
    assert len(set(pix.samples)) > 100


def test_images_to_pdf_does_not_render_other_files_named_as_images(client):
    """O MuPDF, que agora monta o PDF, abre também HTML e PDF: o que não for imagem é recusado antes."""
    for content in (b"<html><body><h1>Oi</h1></body></html>", make_pdf(1)):
        r = post(client, "jpg-to-pdf", ("foto.jpg", content))
        assert r.status_code == 400 and "não é uma imagem válida" in r.get_data(as_text=True)


def test_images_to_pdf_puts_the_jpeg_in_as_it_is(client):
    """A foto era recomprimida (o PDF ficava maior que ela e pior) e perdia o perfil de cor."""
    from PIL import ImageCms
    icc = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    photo = make_photo((400, 300), quality=85)
    buf = io.BytesIO()
    Image.open(io.BytesIO(photo)).save(buf, "JPEG", quality=85, icc_profile=icc)
    r = post(client, "jpg-to-pdf", ("foto.jpg", buf.getvalue()))
    doc = pymupdf.open(stream=r.data)
    xref = doc[0].get_images(full=True)[0][0]
    assert doc.xref_stream_raw(xref) == buf.getvalue()
    assert "ICCBased" in doc.xref_object(xref) + doc.xref_object(int(doc.xref_get_key(xref, "ColorSpace")[1].split()[0]))


needs_libreoffice = pytest.mark.skipif(not euamopdf.find_soffice(), reason="LibreOffice não instalado")


@needs_libreoffice
def test_word_to_pdf(client):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("Olá, documento")
    buf = io.BytesIO()
    document.save(buf)
    r = post(client, "word-to-pdf", ("carta.docx", buf.getvalue()))
    assert r.status_code == 200
    assert "Olá, documento" in texts(r.data)[0]


def tracked_changes_docx():
    """.docx com alterações controladas pendentes, como o de um contrato em negociação."""
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def change(kind, text):
        element = OxmlElement(f"w:{kind}")
        element.set(qn("w:id"), str(len(text)))
        element.set(qn("w:author"), "Revisor")
        run, t = OxmlElement("w:r"), OxmlElement("w:delText" if kind == "del" else "w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        element.append(run)
        return element

    document = docx.Document()
    value = document.add_paragraph("Valor do contrato: ")
    value._p.append(change("del", "R$ 9.000 (VALOR ANTIGO)"))
    value._p.append(change("ins", "R$ 10.000 (VALOR NOVO)"))
    gone = document.add_paragraph()  # parágrafo inteiro excluído, com a marca de parágrafo
    gone._p.append(change("del", "CLÁUSULA EXCLUÍDA"))
    deleted_mark, mark = OxmlElement("w:del"), OxmlElement("w:rPr")
    deleted_mark.set(qn("w:id"), "99")
    deleted_mark.set(qn("w:author"), "Revisor")
    mark.append(deleted_mark)
    gone._p.get_or_add_pPr().append(mark)
    commented = document.add_paragraph("Parágrafo com comentário.")
    if hasattr(document, "add_comment"):  # python-docx 1.2 em diante
        document.add_comment(commented.runs[0], text="COMENTÁRIO DO REVISOR", author="Revisor")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_tracked_changes_are_accepted_before_libreoffice(tmp_path):
    """O LibreOffice imprimia o texto excluído riscado ao lado do inserido; o iLovePDF, que converte
    pelo Word, mostra a versão final."""
    docx = pytest.importorskip("docx")
    src = tmp_path / "contrato.docx"
    src.write_bytes(tracked_changes_docx())
    out = euamopdf.accept_changes(src, tmp_path)
    texts = [p.text for p in docx.Document(out).paragraphs]
    assert texts == ["Valor do contrato: R$ 10.000 (VALOR NOVO)", "Parágrafo com comentário."]
    assert b"w:del" not in unzip(out.read_bytes())["word/document.xml"]


def test_docx_without_tracked_changes_goes_unchanged(tmp_path):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("Sem alterações")
    src = tmp_path / "carta.docx"
    document.save(src)
    assert euamopdf.accept_changes(src, tmp_path) == src


@needs_libreoffice
def test_word_to_pdf_shows_the_final_version(client):
    r = post(client, "word-to-pdf", ("contrato.docx", tracked_changes_docx()))
    pdf = pymupdf.open(stream=r.data)
    text = pdf[0].get_text()
    assert "VALOR NOVO" in text
    assert "VALOR ANTIGO" not in text and "EXCLUÍDA" not in text and "COMENTÁRIO" not in text
    assert not pdf[0].get_drawings()  # sem a barra de alteração na margem


def test_microsoft_word_accepts_the_changes_before_saving(fake_msoffice, tmp_path):
    document = fake_msoffice.Documents.Open.return_value
    order = []
    document.Revisions.AcceptAll.side_effect = lambda: order.append("aceita")
    document.DeleteAllComments.side_effect = lambda: order.append("apaga comentários")
    document.SaveAs.side_effect = lambda *args, **kwargs: order.append("salva o PDF")
    euamopdf.msoffice_to_pdf(tmp_path / "a.docx", tmp_path / "a.pdf", "Word.Application")
    assert order == ["aceita", "apaga comentários", "salva o PDF"]  # só no documento aberto, que não é salvo


@needs_libreoffice
def test_corrupted_office_file_is_reported(client):
    r = post(client, "word-to-pdf", ("quebrado.docx", b"PK\x03\x04lixo"))
    assert r.status_code == 400


# O LibreOffice e o Microsoft Office simulados: estes testes rodam em qualquer sistema, com ou sem eles

@pytest.fixture
def fake_soffice(monkeypatch):
    """Troca o LibreOffice por um que só registra a chamada; devolve a lista de chamadas."""
    calls = []
    monkeypatch.setattr(euamopdf, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(euamopdf, "msoffice_to_pdf", mock.Mock(side_effect=OSError("sem Office")))  # no Windows
    monkeypatch.setattr(euamopdf, "run_soffice", lambda args, timeout: calls.append((args, {"timeout": timeout})))
    return calls


def test_libreoffice_runs_with_its_own_profile_and_a_time_limit(client, fake_soffice):
    post(client, "word-to-pdf", ("a.docx", b"PK"))
    post(client, "word-to-pdf", ("a.docx", b"PK"))
    (first, options), (second, _) = fake_soffice
    assert options["timeout"] == euamopdf.LIBREOFFICE_TIMEOUT
    profiles = [next(a for a in args if a.startswith("-env:UserInstallation=")) for args in (first, second)]
    outdir = Path(first[first.index("--outdir") + 1])
    assert profiles[0] == f"-env:UserInstallation={(outdir / 'perfil').as_uri()}"  # dentro da pasta da conversão
    assert profiles[0] != profiles[1]  # cada conversão com o seu


@pytest.mark.parametrize("run, message", [
    (lambda args, timeout: None, "não conseguiu converter"),  # sai sem erro e sem PDF
    (mock.Mock(side_effect=subprocess.TimeoutExpired("soffice", 180)), "demorou demais"),
])
def test_libreoffice_failures_are_reported(client, fake_soffice, monkeypatch, run, message):
    monkeypatch.setattr(euamopdf, "run_soffice", run)
    r = post(client, "word-to-pdf", ("a.docx", b"PK"))
    assert r.status_code == 400 and message in r.get_data(as_text=True)


@pytest.mark.skipif(not Path("/proc/self/stat").exists(), reason="usa o /proc do Linux para ver os processos")
def test_libreoffice_timeout_kills_every_process_it_started(tmp_path):
    """O soffice só lança o soffice.bin, que é quem converte: o timeout matava só o primeiro, e um
    documento que trava deixava o LibreOffice rodando para sempre."""
    grandchild = tmp_path / "neto.pid"
    soffice = tmp_path / "soffice"
    soffice.write_text(f"#!/bin/sh\nsleep 60 &\necho $! > {grandchild}\nwait\n")
    soffice.chmod(0o755)
    with pytest.raises(subprocess.TimeoutExpired):
        euamopdf.run_soffice([str(soffice)], timeout=1)
    try:
        state = Path(f"/proc/{grandchild.read_text().strip()}/stat").read_text().rsplit(")", 1)[1].split()[0]
    except (FileNotFoundError, ProcessLookupError):  # sumiu antes ou durante a leitura
        state = "morto"
    assert state in ("morto", "Z", "X")  # Z e X: morto, sendo recolhido


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-8", "cp1252"])
@pytest.mark.parametrize("separator", [";", ","])
def test_csv_goes_to_the_converter_as_a_spreadsheet(client, fake_soffice, monkeypatch, encoding, separator):
    """O Excel em português grava CSV com ponto e vírgula, vírgula decimal e em Windows-1252; o
    LibreOffice e o Excel por automação o liam com vírgula, e a tabela saía numa coluna só."""
    openpyxl = pytest.importorskip("openpyxl")
    sheets = []
    monkeypatch.setattr(euamopdf, "run_soffice", lambda args, timeout: sheets.append(openpyxl.load_workbook(args[-1]).active))
    value = "1.234,56" if separator == ";" else '"1.234,56"'
    csv = f"Nome{separator}Cidade{separator}Valor\nJoão{separator}São Paulo{separator}{value}\n"
    post(client, "excel-to-pdf", ("Relatório de vendas.csv", csv.encode(encoding)))
    sheet, = sheets
    assert sheet.title == "Relatório de vendas"  # o LibreOffice o imprime no cabeçalho da página
    assert [[cell.value for cell in row] for row in sheet.iter_rows()] == [
        ["Nome", "Cidade", "Valor"], ["João", "São Paulo", "1.234,56"]]


@needs_libreoffice
def test_brazilian_csv_to_pdf(client):
    csv = "Nome;Cidade;Valor\nJoão;São Paulo;1.234,56\n".encode("cp1252")
    r = post(client, "excel-to-pdf", ("vendas.csv", csv))
    words = [w[4] for w in pymupdf.open(stream=r.data)[0].get_text("words")]
    assert "1.234,56" in words and "João" in words
    assert not any(";" in word for word in words)


def test_libreoffice_is_found_in_program_files_on_windows(monkeypatch, tmp_path):
    exe = tmp_path / "LibreOffice" / "program" / "soffice.exe"
    exe.parent.mkdir(parents=True)
    exe.touch()
    monkeypatch.setattr(euamopdf.sys, "platform", "win32")
    monkeypatch.setattr(euamopdf.shutil, "which", lambda name: None)  # o instalador não põe no PATH
    monkeypatch.setenv("PROGRAMFILES", str(tmp_path))
    assert euamopdf.find_soffice() == str(exe)


@pytest.mark.parametrize("app_name, opened_files, save", [
    ("Word.Application", "Documents", "SaveAs"),
    ("Excel.Application", "Workbooks", "ExportAsFixedFormat"),
    ("PowerPoint.Application", "Presentations", "SaveAs"),
])
def test_microsoft_office_uses_its_own_instance_and_always_quits(monkeypatch, tmp_path, app_name, opened_files, save):
    office = mock.MagicMock()
    document = getattr(office, opened_files).Open.return_value
    getattr(document, save).side_effect = RuntimeError("falhou no meio")
    client_module = types.SimpleNamespace(DispatchEx=mock.Mock(return_value=office))
    pythoncom = mock.Mock()
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com", types.SimpleNamespace(client=client_module))
    monkeypatch.setitem(sys.modules, "win32com.client", client_module)
    with pytest.raises(RuntimeError):
        euamopdf.msoffice_to_pdf(tmp_path / "a", tmp_path / "a.pdf", app_name)
    client_module.DispatchEx.assert_called_once_with(app_name)  # instância própria: não fecha o Office do usuário
    assert getattr(office, opened_files).Open.call_args.kwargs["ReadOnly"] is True
    document.Close.assert_called_once()
    office.Quit.assert_called_once()
    pythoncom.CoUninitialize.assert_called_once()


OFFICE_DOCUMENTS = {"Word.Application": "Documents", "Excel.Application": "Workbooks", "PowerPoint.Application": "Presentations"}


@pytest.fixture
def fake_msoffice(monkeypatch):
    """Microsoft Office simulado: devolve o objeto do aplicativo, que registra o que o app fez nele."""
    office = mock.MagicMock()
    client_module = types.SimpleNamespace(DispatchEx=mock.Mock(return_value=office))
    monkeypatch.setitem(sys.modules, "pythoncom", mock.Mock())
    monkeypatch.setitem(sys.modules, "win32com", types.SimpleNamespace(client=client_module))
    monkeypatch.setitem(sys.modules, "win32com.client", client_module)
    return office


@pytest.mark.parametrize("office_opens, message", [
    (True, "O Microsoft Office não conseguiu converter"),  # abriu e falhou no arquivo: não é falta de programa
    (False, "é preciso ter o Microsoft Office"),
])
def test_office_failure_without_libreoffice_says_what_went_wrong(fake_msoffice, monkeypatch, tmp_path, office_opens, message):
    monkeypatch.setattr(euamopdf.sys, "platform", "win32")
    monkeypatch.setattr(euamopdf, "find_soffice", lambda: None)
    if office_opens:
        fake_msoffice.Documents.Open.side_effect = RuntimeError("arquivo corrompido")
    else:
        sys.modules["win32com.client"].DispatchEx.side_effect = OSError("sem Office")
    (tmp_path / "a.docx").write_bytes(b"PK")
    with pytest.raises(euamopdf.UserError, match=message):
        euamopdf.office_to_pdf(tmp_path / "a.docx", tmp_path, "Word.Application")


@pytest.mark.parametrize("app_name, password", [("Word.Application", "PasswordDocument"), ("Excel.Application", "Password")])
def test_microsoft_office_does_not_wait_for_a_password(fake_msoffice, tmp_path, app_name, password):
    """Com o Office invisível, a janela de senha de um documento protegido ninguém via, e a conversão
    ficava presa para sempre. Com uma senha qualquer, o Office dá erro na hora; sem proteção, ela é
    ignorada."""
    euamopdf.msoffice_to_pdf(tmp_path / "a", tmp_path / "a.pdf", app_name)
    assert getattr(fake_msoffice, OFFICE_DOCUMENTS[app_name]).Open.call_args.kwargs[password]


def test_office_document_with_a_password_is_reported(client, fake_soffice):
    # .docx, .xlsx e .pptx com senha não são ZIP: o pacote vai criptografado num contêiner OLE
    protected = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504) + "EncryptedPackage".encode("utf-16-le")
    r = post(client, "word-to-pdf", ("contrato.docx", protected))
    assert r.status_code == 400 and "tem senha" in r.get_data(as_text=True)
    assert fake_soffice == []  # nem chega ao conversor


@pytest.mark.parametrize("app_name", OFFICE_DOCUMENTS)
def test_microsoft_office_opens_files_with_macros_disabled(fake_msoffice, tmp_path, app_name):
    opened = getattr(fake_msoffice, OFFICE_DOCUMENTS[app_name]).Open
    security_when_opened = []
    opened.side_effect = lambda *args, **kwargs: security_when_opened.append(fake_msoffice.AutomationSecurity) or mock.MagicMock()
    euamopdf.msoffice_to_pdf(tmp_path / "a", tmp_path / "a.pdf", app_name)
    # Por automação, o Office roda as macros do arquivo sem perguntar; 3 (ForceDisable) as desliga antes de abrir
    assert security_when_opened == [3]
    if app_name == "Excel.Application":
        assert opened.call_args.kwargs["UpdateLinks"] == 0  # não busca planilhas vinculadas pela rede
    if app_name == "Word.Application":
        assert opened.call_args.kwargs["AddToRecentFiles"] is False


# --- Validação ---

@pytest.mark.parametrize("action, files, message", [
    ("nao-existe", [("a.pdf", make_pdf())], "desconhecida"),
    ("merge-pdf", [], "Escolha um arquivo"),
    ("split-pdf", [("a.jpg", make_image())], "não é do tipo aceito"),
    ("split-pdf", [("a.pdf", make_pdf()), ("b.pdf", make_pdf())], "um arquivo por vez"),
    ("split-pdf", [("a.pdf", b"")], "vazio"),
    ("split-pdf", [("a.pdf", b"isto nao e um pdf")], "não é um PDF válido"),
    ("jpg-to-pdf", [("a.png", b"isto nao e uma imagem")], "não é uma imagem válida"),
    ("split-pdf", [("a.pdf", make_pdf(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="1", owner_pw="1"))],
     "protegido por senha"),
])
def test_rejects_bad_input_with_message(client, action, files, message):
    r = post(client, action, *files)
    assert r.status_code == 400
    assert message in r.get_data(as_text=True)


def test_unexpected_error_gives_generic_message(client, monkeypatch):
    def broken(files, form, tmp):
        raise RuntimeError("/caminho/interno/que/nao/deve/vazar")
    monkeypatch.setitem(euamopdf.ACTIONS, "merge-pdf", (broken, (".pdf",), True))
    r = post(client, "merge-pdf", ("a.pdf", make_pdf()))
    assert r.status_code == 500
    assert "Não foi possível processar" in r.get_data(as_text=True)
    assert "caminho" not in r.get_data(as_text=True)


def test_unknown_route_is_still_404(client):
    assert client.get("/nao-existe").status_code == 404


def test_rejects_upload_over_limit(client, monkeypatch):
    monkeypatch.setitem(euamopdf.app.config, "MAX_CONTENT_LENGTH", 100)
    assert post(client, "merge-pdf", ("a.pdf", make_pdf())).status_code == 413


def test_parse_pages():
    assert euamopdf.parse_pages("1-3, 5", 5) == [0, 1, 2, 4]
    with pytest.raises(euamopdf.UserError):
        euamopdf.parse_pages("6", 5)


# --- Interface ---

def test_index_has_a_button_for_every_tool(client):
    page = client.get("/").get_data(as_text=True)
    for action in euamopdf.ACTIONS:
        assert f'data-action="{action}"' in page


def test_page_count(client):
    r = client.post("/pages", data={"file": (io.BytesIO(make_pdf(4)), "a.pdf")}, content_type="multipart/form-data")
    assert r.get_json() == {"paginas": 4}
    r = client.post("/pages", data={"file": (io.BytesIO(b"lixo"), "a.pdf")}, content_type="multipart/form-data")
    assert r.status_code == 400


# --- Segurança e isolamento ---

def test_upload_names_cannot_escape_temp_folder(client, tmp_path):
    outside = tmp_path / "fora.pdf"
    relative = Path(tempfile.gettempdir(), "euamopdf_escapou.pdf")
    for name in (str(outside), "../euamopdf_escapou.pdf", "..\\euamopdf_escapou.pdf"):
        assert post(client, "split-pdf", (name, make_pdf())).status_code == 200
    assert not outside.exists()
    assert not relative.exists()


@pytest.mark.parametrize("headers, base_url", [
    ({"Origin": "https://site-malicioso.com"}, "http://127.0.0.1:5000"),
    ({"Origin": "null"}, "http://127.0.0.1:5000"),
    ({}, "http://site-malicioso.com"),  # DNS rebinding
])
def test_rejects_requests_from_other_sites(client, headers, base_url):
    r = client.post("/convert", base_url=base_url, headers=headers, content_type="multipart/form-data",
                    data={"action": "merge-pdf", "file": (io.BytesIO(make_pdf()), "a.pdf")})
    assert r.status_code == 403


def test_accepts_requests_from_the_app_itself(client):
    r = post(client, "merge-pdf", ("a.pdf", make_pdf()), headers={"Origin": "http://127.0.0.1:5000"})
    assert r.status_code == 200


def test_leaves_no_files_behind(client, tmp_path, monkeypatch):
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    post(client, "split-pdf", ("a.pdf", make_pdf()))
    post(client, "pdf-to-jpg", ("a.pdf", make_pdf()))
    post(client, "split-pdf", ("a.pdf", b"invalido"))
    assert list(tmp_path.iterdir()) == []


@pytest.mark.skipif(not Path("/proc/self/fd").is_dir(), reason="usa o /proc do Linux para ver os arquivos abertos")
def test_no_upload_stays_open_when_the_temp_folder_is_deleted(client, monkeypatch):
    """No Windows, arquivo aberto não pode ser apagado: a limpeza da pasta temporária quebraria e
    a mensagem de erro viraria um erro 500. No Linux a limpeza funciona mesmo assim, então o
    teste confere, na hora dela, que nada lá dentro continua aberto."""
    left_open = []

    class CheckedFolder(tempfile.TemporaryDirectory):
        def cleanup(self):
            for fd in Path("/proc/self/fd").iterdir():
                try:
                    target = str(fd.readlink())
                except OSError:
                    continue  # fechado durante a listagem
                if target.startswith(self.name):
                    left_open.append(target)
            super().cleanup()

    monkeypatch.setattr(tempfile, "TemporaryDirectory", CheckedFolder)
    from pdf2docx import Converter
    monkeypatch.setattr(Converter, "convert", lambda self, *args, **kwargs: 1 / 0)  # falha no meio da conversão
    assert post(client, "split-pdf", ("a.pdf", make_pdf()), pages="9").status_code == 400  # erro com o PDF já aberto
    assert post(client, "pdf-to-word", ("a.pdf", make_pdf())).status_code == 500
    assert post(client, "jpg-to-pdf", ("a.jpg", make_photo((200, 150)))).status_code == 200
    assert left_open == []


def test_same_file_name_never_returns_previous_result(client):
    first = post(client, "split-pdf", ("mesmo.pdf", make_pdf(3)))
    second = post(client, "split-pdf", ("mesmo.pdf", make_pdf(2)))
    assert len(unzip(first.data)) == 3
    assert len(unzip(second.data)) == 2


def test_free_port_skips_busy_port():
    import socket
    with socket.socket() as busy:
        busy.bind(("127.0.0.1", 0))
        busy.listen()
        port = busy.getsockname()[1]
        assert euamopdf.free_port(port) != port
