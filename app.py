from flask import Flask, render_template, request, send_file
from werkzeug.exceptions import HTTPException
import base64
import hashlib
import io
import json
import math
import os
import random
import re
import sys
import shutil
import signal
import socket
import subprocess
import tempfile
import threading
import time
import urllib.request
import warnings
import webbrowser # Biblioteca para abrir o navegador
import xml.etree.ElementTree as ET
import zipfile
import zlib
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from threading import Timer # Para atrasar a abertura em 1 segundo
from urllib.parse import quote, urlparse

import numpy as np
import pymupdf
from PIL import Image, ImageChops, ImageOps


app = Flask(__name__)

# --- CONFIGURAÇÕES ---
LOCAL_HOSTS = {'127.0.0.1', 'localhost'}
MAX_UPLOAD_MB = 500
LIBREOFFICE_TIMEOUT = 180  # segundos
app.config['MAX_CONTENT_LENGTH'] = MAX_UPLOAD_MB * 1024 * 1024


class UserError(Exception):
    """Erro causado pela entrada do usuário; a mensagem é mostrada como está."""


# O Pillow recusa imagens acima de ~179 MP (o dobro deste valor), contra arquivos pequenos que
# explodem ao abrir, e as fotos de 200 MP de celular passavam disso. Erro agora só acima de 300 MP;
# a 200 MP, o Remover fundo usa ~1 GB de memória só para a imagem em RGBA.
Image.MAX_IMAGE_PIXELS = 150_000_000
warnings.filterwarnings('ignore', category=Image.DecompressionBombWarning)  # entre 150 e 300 MP: é o esperado
TOO_BIG = "A imagem é grande demais: o limite é de 300 megapixels."


TEXT = {'Content-Type': 'text/plain; charset=utf-8'}  # mensagem de erro, que pode ter o nome do arquivo: nunca HTML


@app.before_request
def only_local_requests():
    """Recusa requisições de outros sites abertos no navegador (CSRF e DNS rebinding)."""
    origin = request.headers.get('Origin')
    if urlparse('//' + request.host).hostname not in LOCAL_HOSTS or \
            (origin and urlparse(origin).hostname not in LOCAL_HOSTS):
        return "Acesso permitido só a partir deste computador.", 403, TEXT


@app.after_request
def security_headers(response):
    """Segunda camada, além da checagem de origem: só os arquivos do próprio app rodam na página
    (as miniaturas e o ícone são data:, o download é blob:), e nenhum site a põe num iframe."""
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Content-Security-Policy'] = "default-src 'self'; img-src 'self' data: blob:; frame-ancestors 'none'"
    return response


def program_file(*parts):
    """Caminho dentro de Program Files, no Windows, onde os instaladores não mexem no PATH."""
    if sys.platform != 'win32':
        return None
    for base in (os.environ.get('PROGRAMFILES'), os.environ.get('PROGRAMFILES(X86)')):
        if base and (path := Path(base, *parts)).exists():
            return str(path)
    return None


def bundled(name):
    """Arquivo embutido no executável gerado pelo PyInstaller (packaging/build.py); None fora dele."""
    if hasattr(sys, '_MEIPASS') and (path := Path(sys._MEIPASS, name)).exists():
        return path


def find_soffice():
    return shutil.which('soffice') or shutil.which('libreoffice') or \
        program_file('LibreOffice', 'program', 'soffice.exe')


def find_tessdata():
    """Pasta de idiomas do Tesseract (a embutida, se houver), ou None se ele não estiver instalado."""
    if path := bundled('tessdata'):
        return str(path)
    try:
        return pymupdf.get_tessdata()
    except Exception:
        return program_file('Tesseract-OCR', 'tessdata')


def msoffice_to_pdf(src, out, app_name):
    """Converte com o Microsoft Office numa instância própria, sem tocar no Office que o usuário tem aberto."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    office = None
    try:
        office = win32com.client.DispatchEx(app_name)
        # Por automação, o Office roda as macros do arquivo sem perguntar e sem o Modo Protegido (a
        # cópia na pasta temporária não tem a marca de "baixado da internet"): 3 = ForceDisable
        office.AutomationSecurity = 3
        if app_name == "Word.Application":
            office.DisplayAlerts = 0  # wdAlertsNone
            # Senha qualquer: num documento protegido o Office dá erro na hora, em vez de abrir uma
            # janela de senha que ninguém vê (o Office está invisível) e prender a conversão para sempre
            doc = office.Documents.Open(str(src), ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False,
                                        PasswordDocument='-', NoEncodingDialog=True)
            try:
                # A versão final, sem alterações controladas nem comentários, como faz o iLovePDF.
                # Só no documento aberto: ele é só leitura e não é salvo
                try:
                    doc.Revisions.AcceptAll()
                    doc.DeleteAllComments()
                except Exception:
                    app.logger.info("Documento protegido contra edição: vai como está", exc_info=True)
                doc.SaveAs(str(out), FileFormat=17)  # wdFormatPDF
            finally: doc.Close(False)
        elif app_name == "Excel.Application":
            office.DisplayAlerts = False
            wb = office.Workbooks.Open(str(src), ReadOnly=True, UpdateLinks=0,  # sem buscar planilhas vinculadas
                                       Password='-', IgnoreReadOnlyRecommended=True)
            try: wb.ExportAsFixedFormat(0, str(out))  # xlTypePDF
            finally: wb.Close(False)
        else:
            pres = office.Presentations.Open(str(src), ReadOnly=True, WithWindow=False)
            try: pres.SaveAs(str(out), 32)  # ppSaveAsPDF
            finally: pres.Close()
    except Exception as e:
        e.office_opened = office is not None  # abriu e falhou no arquivo, ou nem abriu
        raise
    finally:
        if office is not None:
            office.Quit()
        pythoncom.CoUninitialize()


WORD_PARTS = re.compile(r'word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml')


def accept_changes(src, tmp):
    """Cópia do .docx com as alterações controladas aceitas, ou o próprio src se não houver nenhuma.
    O LibreOffice imprimia o texto excluído riscado ao lado do inserido, com uma barra na margem; o
    iLovePDF, que converte pelo Word, mostra só a versão final."""
    from lxml import etree  # vem com o python-docx
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

    def unwrap(element):
        parent = element.getparent()
        for i, child in enumerate(list(element)):
            parent.insert(parent.index(element) + i, child)
        parent.remove(element)

    parts = {}
    if not zipfile.is_zipfile(src):
        return src  # corrompido: o LibreOffice diz que não conseguiu converter
    with zipfile.ZipFile(src) as z:
        for name in z.namelist():
            data = z.read(name)
            if not WORD_PARTS.fullmatch(name) or not re.search(rb'<w:(del|ins|moveFrom|moveTo)\b|PrChange\b', data):
                continue
            try:
                root = etree.fromstring(data)
            except etree.XMLSyntaxError:
                return src
            for mark in list(root.iter(f'{w}del', f'{w}moveFrom')):
                if mark.getparent() is None:
                    continue  # já saiu junto com o que a continha
                holder = mark.getparent()
                if holder.tag == f'{w}trPr':  # linha de tabela excluída
                    row = holder.getparent()
                    row.getparent().remove(row)
                elif holder.tag == f'{w}rPr' and holder.getparent().tag == f'{w}pPr':
                    # marca de parágrafo excluída: o que sobrou do parágrafo se junta ao seguinte
                    paragraph = holder.getparent().getparent()
                    holder.remove(mark)
                    following = paragraph.getnext()
                    if following is not None and following.tag == f'{w}p':
                        content = [child for child in paragraph if child.tag != f'{w}pPr']
                        start = 1 if following.find(f'{w}pPr') is not None else 0
                        for i, child in enumerate(content):
                            following.insert(start + i, child)
                        paragraph.getparent().remove(paragraph)
                else:
                    holder.remove(mark)
            for mark in list(root.iter(f'{w}ins', f'{w}moveTo')):
                if mark.getparent().tag in (f'{w}rPr', f'{w}trPr'):
                    mark.getparent().remove(mark)  # só a marca de inserido
                else:
                    unwrap(mark)
            for change in list(root.iter('{*}*')):
                if change.tag.endswith('PrChange') or change.tag.split('}')[-1].startswith(('moveFromRange', 'moveToRange')):
                    change.getparent().remove(change)  # formatação antiga: fica a nova
            parts[name] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        if not parts:
            return src
        out = tmp / 'final.docx'
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as final:
            for item in z.infolist():
                final.writestr(item, parts.get(item.filename) or z.read(item))
    return out


def run_soffice(args, timeout):
    """Roda o LibreOffice num grupo de processos próprio. O soffice só lança o soffice.bin, que é quem
    converte: no tempo esgotado, matar só o primeiro deixava o outro rodando para sempre (e, no
    Windows, segurando os arquivos da pasta temporária). Mata o grupo inteiro, como o build.py."""
    windows = sys.platform == 'win32'
    group = {'creationflags': subprocess.CREATE_NEW_PROCESS_GROUP} if windows else {'start_new_session': True}
    proc = subprocess.Popen(args, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **group)
    try:
        proc.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        if windows:
            subprocess.run(['taskkill', '/F', '/T', '/PID', str(proc.pid)], capture_output=True)
        else:
            os.killpg(proc.pid, signal.SIGKILL)
        proc.wait()
        raise


def libreoffice_to_pdf(src, out_dir):
    soffice = find_soffice()
    if not soffice:
        raise UserError("Para converter arquivos do Office é preciso ter o Microsoft Office (Windows) ou o LibreOffice instalado.")
    # Perfil próprio por conversão: não conflita com outro LibreOffice aberto nem com conversões simultâneas
    profile = (out_dir / 'perfil').as_uri()
    try:
        run_soffice([soffice, '--headless', '--norestore', f'-env:UserInstallation={profile}',
                     '--convert-to', 'pdf', '--outdir', str(out_dir), str(src)], timeout=LIBREOFFICE_TIMEOUT)
    except subprocess.TimeoutExpired:
        raise UserError("O LibreOffice demorou demais para converter este arquivo.")
    # O LibreOffice sai com código 0 mesmo quando falha: o que vale é o PDF existir
    out = out_dir / f"{src.stem}.pdf"
    if not out.exists():
        raise UserError("O LibreOffice não conseguiu converter este arquivo. Veja se ele abre normalmente.")
    return out


def office_to_pdf(src, tmp, app_name):
    # .docx, .xlsx e .pptx com senha não são ZIP: o pacote vai criptografado num contêiner OLE
    with open(src, 'rb') as f:
        head = f.read(65536)
    if head.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1') and 'EncryptedPackage'.encode('utf-16-le') in head:
        raise UserError("Este documento tem senha: remova-a no Office antes de converter.")
    if sys.platform == 'win32':
        out = tmp / 'saida.pdf'
        try:
            msoffice_to_pdf(src, out, app_name)
            return out
        except Exception as e:
            app.logger.info("O Microsoft Office não converteu; tentando o LibreOffice", exc_info=True)
            if getattr(e, 'office_opened', False) and not find_soffice():
                raise UserError("O Microsoft Office não conseguiu converter este arquivo. Veja se ele abre normalmente.")
    out_dir = tmp / 'libreoffice'
    out_dir.mkdir()
    if src.suffix == '.docx':  # .doc, .odt e .rtf com alterações saem com a marcação (decisão D2)
        src = accept_changes(src, tmp)
    return libreoffice_to_pdf(src, out_dir)


def open_pdf(source, password=None):
    """Abre um PDF a partir do caminho ou dos bytes do arquivo, com a senha se ele pedir uma."""
    try:
        # Abre da memória: no Windows, um arquivo que fica aberto impede apagar a pasta temporária
        data = source if isinstance(source, bytes) else Path(source).read_bytes()
        doc = pymupdf.open(stream=data, filetype='pdf')
    except Exception:
        raise UserError("O arquivo não é um PDF válido ou está corrompido.")
    if doc.needs_pass:
        if password is None:
            raise UserError("Este PDF está protegido por senha. Use \"Remover senha\" antes.")
        if not doc.authenticate(password):
            raise UserError("Senha incorreta.")
    return doc


def pdf_bytes(doc):
    return doc.tobytes(garbage=3, deflate=True)


def parse_pages(spec, count):
    """Converte '1-3, 5' em índices [0, 1, 2, 4], na ordem em que foram escritos."""
    pages = []
    for part in spec.replace(' ', '').split(','):
        if not part:
            continue
        m = re.fullmatch(r'(\d+)(?:-(\d+))?', part)
        if not m:
            raise UserError(f"Não entendi \"{part}\". Escreva as páginas assim: 1-3, 5, 8.")
        start, end = int(m[1]), int(m[2] or m[1])
        if start > end:
            raise UserError(f"Intervalo invertido: \"{part}\". Use o menor número primeiro.")
        if start < 1 or end > count:
            raise UserError(f"Página fora do documento: \"{part}\". O PDF tem {count} página(s).")
        pages.extend(range(start - 1, end))
    if not pages:
        raise UserError("Informe ao menos uma página.")
    return pages


# --- AÇÕES ---
# Cada ação recebe a lista de (caminho salvo, nome original sem extensão),
# o formulário e a pasta temporária da requisição, e devolve (bytes, nome do download)
# e, opcionalmente, uma mensagem para a interface mostrar.

def merge_pdf(files, form, tmp):
    out = pymupdf.open()
    for path, _ in files:
        out.insert_pdf(open_pdf(path))
    return pdf_bytes(out), "PDF_Unido.pdf"


def split_pdf(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path)
    spec = form.get('pages', '').strip()

    if spec:
        doc.select(parse_pages(spec, doc.page_count))
        return pdf_bytes(doc), f"{base}_paginas.pdf"

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for i in range(doc.page_count):
            page = pymupdf.open()
            page.insert_pdf(doc, from_page=i, to_page=i)
            z.writestr(f"pag_{i+1}.pdf", pdf_bytes(page))
    return buf.getvalue(), f"{base}_dividido.zip"


def csv_to_xlsx(path, tmp, title):
    """Planilha a partir do CSV. O Office e o LibreOffice leem CSV com vírgula, e o do Excel em
    português (ponto e vírgula, vírgula decimal, Windows-1252) saía numa coluna só, com os centavos
    numa coluna à parte e sem os acentos."""
    import csv
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    raw = path.read_bytes()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('cp1252', errors='replace')
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=',;\t')
    except csv.Error:  # uma coluna só, ou não deu para decidir
        dialect = csv.excel
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\\*?:/\[\]]', '', title)[:31] or 'Planilha'  # o nome que vai no cabeçalho
    for row in csv.reader(io.StringIO(text), dialect):
        ws.append([ILLEGAL_CHARACTERS_RE.sub('', cell) for cell in row])
    out = tmp / 'planilha.xlsx'
    wb.save(out)
    return out


def office_action(app_name):
    def convert(files, form, tmp):
        path, base = files[0]
        if path.suffix == '.csv':
            path = csv_to_xlsx(path, tmp, base)
        return office_to_pdf(path, tmp, app_name).read_bytes(), f"{base}.pdf"
    return convert


A4 = pymupdf.paper_rect('a4')


def has_alpha(img):
    return img.mode in ('RGBA', 'LA', 'PA') or 'transparency' in img.info


def on_white(img):
    """Imagem em RGB, com a transparência sobre fundo branco."""
    if not has_alpha(img):
        return img.convert('RGB')
    img = img.convert('RGBA')
    bg = Image.new('RGB', img.size, 'white')
    bg.paste(img, mask=img.getchannel('A'))
    return bg


def open_image(path, reduce_to=None):
    """Abre a imagem já na orientação da foto: (imagem, formato original). Com reduce_to, um JPEG
    grande é decodificado direto num tamanho menor (mas ainda o dobro de reduce_to): bem mais rápido."""
    try:
        with Image.open(path) as original:
            fmt = 'JPEG' if original.format == 'MPO' else original.format  # MPO: JPEG de alguns celulares
            if reduce_to:
                original.draft(original.mode, (reduce_to * 2, reduce_to * 2))
            img = ImageOps.exif_transpose(original)
            img.load()
    except Image.DecompressionBombError:
        raise UserError(TOO_BIG)
    except Exception:
        raise UserError("O arquivo não é uma imagem válida.")
    return img, fmt


def load_image(path):
    """Imagem pronta para o PDF: na orientação da foto e com a transparência sobre fundo branco."""
    img, fmt = open_image(path)
    img = on_white(img)
    buf = io.BytesIO()
    if fmt == 'JPEG':  # foto continua JPEG; o resto vira PNG para não borrar texto
        img.save(buf, 'JPEG', quality=95)
    else:
        img.save(buf, 'PNG')
    return buf.getvalue(), img.size


MUPDF_IMAGES = {'JPEG': 'jpg', 'PNG': 'png', 'TIFF': 'tif', 'GIF': 'gif', 'BMP': 'bmp'}


def image_pages(path):
    """A imagem como PDF, uma página por quadro (os TIFFs de scanner têm vários). O MuPDF põe o JPEG
    como está, sem recomprimir e com o perfil de cor, segue a orientação da foto e lê PNG de 16 bits;
    o resto (WebP, o MPO de alguns celulares) vai pelo Pillow."""
    try:
        with Image.open(path) as img:  # só o cabeçalho: o MuPDF abriria até HTML com nome de imagem
            fmt = img.format
    except Image.DecompressionBombError:
        raise UserError(TOO_BIG)
    except Exception:
        raise UserError("O arquivo não é uma imagem válida.")
    try:
        # Da memória: no Windows, um arquivo que fica aberto impede apagar a pasta temporária
        return pymupdf.open('pdf', pymupdf.open(stream=path.read_bytes(), filetype=MUPDF_IMAGES[fmt]).convert_to_pdf())
    except Exception:
        data, (w, h) = load_image(path)
        doc = pymupdf.open()
        doc.new_page(width=w, height=h).insert_image(pymupdf.Rect(0, 0, w, h), stream=data)
        return doc


def images_to_pdf(files, form, tmp):
    out = pymupdf.open()
    for path, _ in files:
        images = image_pages(path)
        for image in images:
            # Página A4 em pé ou deitada, conforme a imagem, com a imagem inteira centralizada
            w, h = image.rect.width, image.rect.height
            size = (A4.width, A4.height) if h >= w else (A4.height, A4.width)
            page = out.new_page(width=size[0], height=size[1])
            page.show_pdf_page(page.rect, images, image.number)
    name = f"{files[0][1]}.pdf" if len(files) == 1 else "Imagens.pdf"
    return pdf_bytes(out), name


# Códigos da fonte ZapfDingbats que os checkboxes e botões de opção do PDF usam, no Unicode
DINGBATS = str.maketrans({'3': '✓', '4': '✔', '5': '✕', '6': '✖', '7': '✗', '8': '✘', 'l': '●', 'n': '■', 'u': '◆', 'H': '★'})


def fix_docx(path):
    """Corrige o .docx do pdf2docx. Ele grava a largura certa em cada célula, mas monta a grade da
    tabela com as colunas iguais, e os editores desenham pela grade: a coluna estreitada quebrava o
    texto, e a altura exata da linha cortava o resto. A grade passa a vir de uma linha sem mesclagem.
    E grava a marca do checkbox como o código dela na ZapfDingbats ("3"), que nenhum editor tem."""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(path)
    for fonts in doc.element.body.iter(qn('w:rFonts')):
        if 'dingbats' in (fonts.get(qn('w:ascii')) or '').lower():
            for text in fonts.getparent().getparent().iter(qn('w:t')):  # rFonts > rPr > r
                text.text = (text.text or '').translate(DINGBATS)
            for attr in ('w:ascii', 'w:hAnsi'):
                fonts.set(qn(attr), 'Segoe UI Symbol')
    for table in doc.element.body.iter(qn('w:tbl')):  # também as tabelas dentro de células
        grid = table.find(qn('w:tblGrid')).findall(qn('w:gridCol'))
        for row in table.findall(qn('w:tr')):
            widths = [tc.find(f"{qn('w:tcPr')}/{qn('w:tcW')}") for tc in row.findall(qn('w:tc'))]
            spans = row.findall(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
            if len(widths) == len(grid) and not spans and None not in widths:
                for col, width in zip(grid, widths):
                    col.set(qn('w:w'), width.get(qn('w:w')))
                break
    doc.save(path)


WORD_MIN_TEXT = 0.97  # parte do texto do PDF que o .docx precisa ter; abaixo disso, outra tentativa


def text_chars(texts):
    """Caracteres do texto, sem os espaços: o pdf2docx às vezes junta palavras ("écelebrado")."""
    return Counter(c for text in texts for c in text if not c.isspace())


def docx_text_share(pdf_chars, path):
    """Quanto do texto do PDF está no .docx, de 0 a 1."""
    from docx import Document
    from docx.oxml.ns import qn
    got = text_chars(t.text or '' for t in Document(path).element.body.iter(qn('w:t')))
    total = sum(pdf_chars.values())
    return sum(min(n, got[c]) for c, n in pdf_chars.items()) / total if total else 1


FONT_STYLE = r'(?:[\s-]+(?:Regular|Bold|Italic|Oblique|Light|Medium|Semibold|Black))+$'


def font_objects(doc, xref):
    """A fonte e, quando estão em objetos separados, a descendente (Type0) e os descritores."""
    found = [xref]
    kind, value = doc.xref_get_key(xref, 'DescendantFonts')
    if kind == 'xref':
        value = doc.xref_object(int(value.split()[0]), compressed=True)
    if kind != 'null':
        found += [int(n) for n in re.findall(r'^\[\s*(\d+) 0 R\s*\]$', value)]
    for font in list(found):
        kind, value = doc.xref_get_key(font, 'FontDescriptor')
        if kind == 'xref':
            found.append(int(value.split()[0]))
    return found


def name_fonts(doc):
    """Dá às fontes de nome genérico o nome da família gravado no arquivo da fonte embutida e devolve
    quantas mudaram. O "Microsoft Print to PDF" chama todas de CIDFont+F1, F2...: o pdf2docx acha a
    fonte pelo nome, e o editor não conhecia o genérico e usava outra, mais larga, que cortava o texto."""
    def simple(name):
        return re.sub(r'[\W_]', '', name).lower()
    seen, renamed = set(), 0
    for page in doc:
        for xref, _, _, basefont, *_ in page.get_fonts(full=True):
            if xref in seen:
                continue
            seen.add(xref)
            buffer = doc.extract_font(xref)[3]
            try:
                family = re.sub(FONT_STYLE, '', pymupdf.Font(fontbuffer=buffer).name) if buffer else ''
            except Exception:  # fonte embutida que o MuPDF não lê: fica como está
                continue
            if not family or simple(family) in simple(basefont.split('+')[-1]):
                continue
            # O nome aparece em BaseFont e FontName, também nos dicionários escritos dentro da fonte
            old = re.compile('/' + re.escape(basefont) + r'(?=[\s/<>\[\]()]|$)')
            for obj in font_objects(doc, xref):
                source = doc.xref_object(obj, compressed=True)
                if old.search(source):
                    doc.update_object(obj, old.sub('/' + family.replace(' ', '#20'), source))
            renamed += 1
    return renamed


def pdf_to_word(files, form, tmp):
    path, base = files[0]
    pdf = open_pdf(path)
    # Anotações e campos (o selo visível de uma assinatura, checkboxes) viram conteúdo da página: o
    # pdf2docx só tira as imagens do conteúdo. A assinatura deixa de valer só nesta cópia
    annotated = any(page.first_annot or page.first_widget for page in pdf)
    if annotated:
        pdf.bake()
    pdf_chars = text_chars(page.get_text() for page in pdf)
    if name_fonts(pdf) or annotated:
        path = tmp / 'entrada_word.pdf'
        pdf.save(path)
    from pdf2docx import Converter  # importação lenta: só quando usada

    def convert(out, **options):
        cv = Converter(str(path))
        try:
            cv.convert(str(out), **options)
        finally:
            cv.close()
        return docx_text_share(pdf_chars, out), out

    # A detecção de tabelas sem bordas do pdf2docx às vezes descarta o texto de tabelas com bordas
    # (um PDF do "Microsoft Print to PDF" perdia 41% das palavras). Sem ela o texto fica, mas tabelas
    # lado a lado saem uma embaixo da outra: por isso ela só é desligada quando falta texto.
    share, out = convert(tmp / 'saida.docx')
    if share < WORD_MIN_TEXT:
        share, out = max((share, out), convert(tmp / 'saida2.docx', parse_stream_table=False), key=lambda r: r[0])
    fix_docx(out)
    lost = ["Parte do texto do PDF pode não ter sido convertida: confira o documento."] if share < WORD_MIN_TEXT else []
    return out.read_bytes(), f"{base}.docx", *lost


def pdf_to_jpg(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path)
    dpi = {'72': 72, '300': 300}.get(form.get('dpi'), 150)

    def jpg(page):
        return page.get_pixmap(dpi=dpi).tobytes('jpg', jpg_quality=90)

    if doc.page_count == 1:
        return jpg(doc[0]), f"{base}.jpg"
    digits = len(str(doc.page_count))  # pag_01, pag_02... para ordenar certo
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as z:  # JPEG já vem comprimido
        for page in doc:
            z.writestr(f"{base}_pag_{page.number + 1:0{digits}d}.jpg", jpg(page))
    return buf.getvalue(), f"{base}_imagens.zip"


AMOUNT = re.compile(r'(-?)\s*(R\$)?\s*(-?)\s*(\d{1,3}(?:\.\d{3})+|\d+),(\d+)')
DATE = re.compile(r'(\d\d)/(\d\d)/(\d{4})')


def sheet_cell(ws, text):
    """Célula da planilha: número para os valores com vírgula decimal (-1.234,56, R$ 10,00) e data
    para dd/mm/aaaa, para o Excel somar e ordenar. O resto, inteiros inclusive, fica texto: contas,
    CEPs e documentos têm zeros à esquerda."""
    from openpyxl.cell.cell import Cell, ILLEGAL_CHARACTERS_RE
    if not text:
        return None
    if m := AMOUNT.fullmatch(text.strip()):
        sign = -1 if m[1] or m[3] else 1
        cell = Cell(ws, value=sign * float(f"{m[4].replace('.', '')}.{m[5]}"))
        cell.number_format = ('"R$" ' if m[2] else '') + '#,##0.' + '0' * len(m[5])
        return cell
    if m := DATE.fullmatch(text.strip()):
        try:
            cell = Cell(ws, value=datetime(int(m[3]), int(m[2]), int(m[1])))
        except ValueError:  # 31/02/2026 e afins: fica como veio
            return ILLEGAL_CHARACTERS_RE.sub('', text)
        cell.number_format = 'dd/mm/yyyy'
        return cell
    return ILLEGAL_CHARACTERS_RE.sub('', text)


def table_rows(doc):
    """[(página, [(linhas, retângulo) de cada tabela])] e se as tabelas vieram sem linhas desenhadas. Extratos de banco
    só alinham as colunas: quando nenhuma página tem tabela com linhas, a busca é pelo alinhamento do
    texto, que também pega título e rodapé (e pode tomar texto corrido por tabela). Por isso ela só
    entra quando a outra não acha nada, e sem as linhas e colunas vazias que ela cria."""
    found = [(page, [(t.extract(), t.bbox) for t in page.find_tables().tables]) for page in doc]
    if any(tables for _, tables in found):
        return found, False
    found = []
    for page in doc:
        tables = []
        for table in page.find_tables(strategy='text').tables:
            rows = [row for row in table.extract() if any(cell and cell.strip() for cell in row)]
            keep = [i for i in range(len(rows[0]) if rows else 0) if any(row[i] and row[i].strip() for row in rows)]
            if len(rows) >= 2 and len(keep) >= 2:
                tables.append(([[row[i] for i in keep] for row in rows], table.bbox))
        found.append((page, tables))
    return found, True


def pdf_to_excel(files, form, tmp):
    from openpyxl import Workbook
    path, base = files[0]
    doc = open_pdf(path)
    wb = Workbook()
    wb.remove(wb.active)
    pages, borderless = table_rows(doc)
    last = None  # a última tabela da página anterior: (aba, cabeçalho, se termina embaixo)
    for page, tables in pages:
        height = page.rect.height
        for n, (rows, bbox) in enumerate(tables, 1):
            # Tabela que passa para a página seguinte: a primeira desta, no alto, com as mesmas colunas
            # da anterior, que terminou embaixo, entra na mesma aba, sem repetir o cabeçalho
            if n == 1 and last and last[2] and bbox[1] < 0.4 * height and len(rows[0]) == len(last[1]):
                ws, header = last[0], last[1]
                rows = rows[1:] if rows[0] == header else rows
            else:
                ws, header = wb.create_sheet(f"Pág {page.number + 1} - Tabela {n}"), rows[0]
            for row in rows:
                ws.append([sheet_cell(ws, cell) for cell in row])
            if n == len(tables):
                last = (ws, header, bbox[3] > 0.6 * height)
        if not tables:
            last = None
    if not wb.sheetnames:
        scanned = not any(page.get_text().strip() for page in doc)
        raise UserError("Nenhuma tabela encontrada neste PDF." + (" Ele parece escaneado: passe o OCR antes." if scanned else ""))
    buf = io.BytesIO()
    wb.save(buf)
    found = ["O PDF não tem tabelas com linhas: elas foram montadas pelo alinhamento do texto. Confira a planilha."] if borderless else []
    return buf.getvalue(), f"{base}.xlsx", *found


def pdf_to_ppt(files, form, tmp):
    """Cada página vira um slide com a imagem da página (o texto não fica editável)."""
    from pptx import Presentation
    from pptx.util import Pt
    path, base = files[0]
    doc = open_pdf(path)
    prs = Presentation()
    first = doc[0].rect
    scale = min(1, 4000 / max(first.width, first.height))  # o PowerPoint limita o slide a 56 polegadas
    prs.slide_width, prs.slide_height = Pt(first.width * scale), Pt(first.height * scale)
    for page in doc:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # layout em branco
        # Encaixa a página no slide sem distorcer, caso ela tenha outro formato
        fit = min(prs.slide_width / page.rect.width, prs.slide_height / page.rect.height)
        w, h = round(page.rect.width * fit), round(page.rect.height * fit)
        image = io.BytesIO(page.get_pixmap(dpi=150).tobytes('jpg', jpg_quality=92))
        slide.shapes.add_picture(image, (prs.slide_width - w) // 2, (prs.slide_height - h) // 2, w, h)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), f"{base}.pptx"


def organize_pdf(files, form, tmp):
    """Monta o PDF na ordem que o usuário escolheu nas miniaturas: [{"pagina": 3, "giro": 90}, ...]."""
    path, base = files[0]
    doc = open_pdf(path)
    try:
        order = [(int(item['pagina']), int(item.get('giro', 0))) for item in json.loads(form.get('order', ''))]
    except (ValueError, TypeError, KeyError, AttributeError):
        raise UserError("Não entendi a nova ordem das páginas.")
    if not order:
        raise UserError("Deixe ao menos uma página.")
    out = pymupdf.open()
    for number, turn in order:
        if not 1 <= number <= doc.page_count or turn % 90:
            raise UserError("Não entendi a nova ordem das páginas.")
        out.insert_pdf(doc, from_page=number - 1, to_page=number - 1)
        out[-1].set_rotation((out[-1].rotation + turn) % 360)
    return pdf_bytes(out), f"{base}_organizado.pdf"


def thumbnail(page, width=150):
    zoom = width / page.rect.width
    jpg = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).tobytes('jpg', jpg_quality=70)
    return 'data:image/jpeg;base64,' + base64.b64encode(jpg).decode()


def rotate_pdf(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path)
    angle = {'90': 90, '180': 180, '270': 270}.get(form.get('angle'))
    if angle is None:
        raise UserError("Escolha para que lado girar.")
    spec = form.get('pages', '').strip()
    for i in set(parse_pages(spec, doc.page_count)) if spec else range(doc.page_count):
        doc[i].set_rotation((doc[i].rotation + angle) % 360)
    return pdf_bytes(doc), f"{base}_girado.pdf"


def protect_pdf(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path)
    password = form.get('password', '')
    if not password:
        raise UserError("Escolha uma senha.")
    if password != form.get('password2'):
        raise UserError("As senhas digitadas não são iguais.")
    data = doc.tobytes(garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                       user_pw=password, owner_pw=password)
    return data, f"{base}_protegido.pdf"


def unlock_pdf(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path, password=form.get('password', ''))
    # Também remove restrições (impressão, cópia...) de PDFs que abrem sem senha
    if not doc.metadata.get('encryption'):
        raise UserError("Este PDF não tem senha nem restrições para remover.")
    return doc.tobytes(garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_NONE), f"{base}_sem_senha.pdf"


def check_text(text, what):
    """As fontes padrão do PDF só têm os caracteres do Windows-1252 (acentos do português incluídos)."""
    if not text:
        raise UserError(f"Escreva o texto {what}.")
    try:
        text.encode('cp1252')
    except UnicodeEncodeError:
        raise UserError("Use só letras, números, acentos e pontuação comuns (sem emojis ou símbolos especiais).")


def watermark_pdf(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path)
    text = form.get('text', '').strip()
    check_text(text, "da marca d'água")
    for page in doc:
        r = page.rect  # área visível, já considerando a rotação da página
        size = min(100, 0.75 * abs(r.br - r.tl) / pymupdf.get_text_length(text, 'helv', 1))
        center = pymupdf.Point(r.width / 2, r.height / 2) * page.derotation_matrix
        start = pymupdf.Point(center.x - pymupdf.get_text_length(text, 'helv', size) / 2, center.y + size * 0.35)
        # Diagonal subindo da esquerda para a direita, qualquer que seja a rotação da página
        page.insert_text(start, text, fontsize=size, fontname='helv', color=(0.5, 0.5, 0.5),
                         fill_opacity=0.3, morph=(center, pymupdf.Matrix(45 + page.rotation)))
    return pdf_bytes(doc), f"{base}_marca_dagua.pdf"


NUMBER_FORMATS = {
    'n': "{n}",
    'n-de-t': "{n} / {t}",
    'pagina': "Página {n} de {t}",
}


# --- METADADOS ---
# Um PDF guarda os metadados em dois lugares: o dicionário Info e o pacote XMP. Leitores como o
# do Firefox e o Acrobat dão preferência ao XMP, então os dois são atualizados juntos.

META_FIELDS = {  # chave no PyMuPDF: campo do formulário
    'title': 'meta_title', 'author': 'meta_author', 'subject': 'meta_subject', 'keywords': 'meta_keywords',
    'creator': 'meta_creator', 'producer': 'meta_producer',
}
META_DATES = {'creationDate': 'meta_created', 'modDate': 'meta_modified'}
XMP_NS = {
    'x': 'adobe:ns:meta/', 'rdf': 'http://www.w3.org/1999/02/22-rdf-syntax-ns#',
    'dc': 'http://purl.org/dc/elements/1.1/', 'pdf': 'http://ns.adobe.com/pdf/1.3/',
    'xmp': 'http://ns.adobe.com/xap/1.0/', 'xmpMM': 'http://ns.adobe.com/xap/1.0/mm/',
    'pdfaid': 'http://www.aiim.org/pdfa/ns/id/', 'exif': 'http://ns.adobe.com/exif/1.0/',
    'photoshop': 'http://ns.adobe.com/photoshop/1.0/',
}
XMP_PROPERTIES = {  # chave no PyMuPDF: (prefixo, propriedade XMP, forma do valor)
    'title': ('dc', 'title', 'Alt'), 'author': ('dc', 'creator', 'Seq'), 'subject': ('dc', 'description', 'Alt'),
    'keywords': ('pdf', 'Keywords', 'text'), 'creator': ('xmp', 'CreatorTool', 'text'),
    'producer': ('pdf', 'Producer', 'text'), 'creationDate': ('xmp', 'CreateDate', 'text'),
    'modDate': ('xmp', 'ModifyDate', 'text'),
}


def parse_pdf_date(value):
    """'D:20260929143000-03'00'' em datetime; None se vazia ou ilegível. Aceita também o que alguns
    programas gravam no lugar: sem o 'D:', ISO ou por extenso."""
    m = re.match(r"(?:D:\s*|(?=\d{8}))(\d{4})(\d\d)?(\d\d)?(\d\d)?(\d\d)?(\d\d)?(Z|[+-]\d\d'?\d\d'?)?", value or '')
    if not m:
        return text_date(value or '')
    year, month, day, hour, minute, second, zone = m.groups()
    tz = None
    if zone == 'Z':
        tz = timezone.utc
    elif zone:
        digits = re.sub(r'\D', '', zone)
        tz = timezone((1 if zone[0] == '+' else -1) * timedelta(hours=int(digits[:2]), minutes=int(digits[2:4])))
    try:
        return datetime(int(year), int(month or 1), int(day or 1), int(hour or 0), int(minute or 0), int(second or 0), tzinfo=tz)
    except ValueError:
        return None


def form_date(value):
    """Data do formulário ('2026-09-29T14:30'), no fuso deste computador; None se vazia."""
    if not value:
        return None
    try:
        return datetime.fromisoformat(value).astimezone()
    except ValueError:
        raise UserError(f"Data inválida: \"{value}\".")


def pdf_date(moment):
    offset = moment.strftime('%z')  # -0300
    return moment.strftime('D:%Y%m%d%H%M%S') + f"{offset[:3]}'{offset[3:]}'"


def update_xmp(xmp, values, properties=XMP_PROPERTIES, drop=None):
    """Troca no XMP as propriedades editadas e mantém o resto (como a identificação PDF/A).
    values: chave em properties -> texto já no formato do XMP ('' apaga); drop(tag): se a
    propriedade deve sair, qualquer que seja. None se o XMP for ilegível."""
    for prefix, uri in XMP_NS.items():
        ET.register_namespace(prefix, uri)
    try:
        root = ET.fromstring(xmp)
    except ET.ParseError:
        return None
    rdf = XMP_NS['rdf']
    descriptions = list(root.iter(f'{{{rdf}}}Description'))
    if not descriptions:
        return None
    if drop:
        for description in descriptions:
            for tag in [t for t in description.attrib if drop(t)]:
                del description.attrib[tag]
            for old in [child for child in description if drop(child.tag)]:
                description.remove(old)
    for key, value in values.items():
        prefix, name, form = properties[key]
        tag = f'{{{XMP_NS[prefix]}}}{name}'
        for description in descriptions:  # a propriedade pode vir como atributo ou como elemento
            description.attrib.pop(tag, None)
            for old in description.findall(tag):
                description.remove(old)
        if not value:
            continue
        prop = ET.SubElement(descriptions[0], tag)
        if form == 'text':
            prop.text = value
            continue
        container = ET.SubElement(prop, f'{{{rdf}}}{form}')
        for text in [v.strip() for v in value.split(',') if v.strip()] if form == 'Bag' else [value]:
            item = ET.SubElement(container, f'{{{rdf}}}li')
            if form == 'Alt':
                item.set('{http://www.w3.org/XML/1998/namespace}lang', 'x-default')
            item.text = text
    body = ET.tostring(root, encoding='unicode')
    return f'<?xpacket begin="\ufeff" id="W5M0MpCehiHzreSzNTczkc9d"?>\n{body}\n<?xpacket end="w"?>'


def read_metadata(doc):
    """Metadados atuais, com os nomes dos campos do formulário. O que faltar no Info vem do XMP:
    há programas que só gravam nele, como pede o PDF 2.0."""
    meta = doc.metadata or {}
    xmp = xmp_values(doc.get_xml_metadata(), XMP_PROPERTIES)
    values = {field: meta.get(key) or xmp.get(key, '') for key, field in META_FIELDS.items()}
    for key, field in META_DATES.items():
        moment = parse_pdf_date(meta.get(key)) or text_date(xmp.get(key, ''))
        if moment and moment.tzinfo:
            moment = moment.astimezone()  # mostra no fuso deste computador
        values[field] = moment.strftime('%Y-%m-%dT%H:%M:%S') if moment else ''
    return values


def edit_metadata(files, form, tmp):
    path, base = files[0]
    if path.suffix != '.pdf':
        return edit_image_metadata(path, base, form)
    doc = open_pdf(path)
    if form.get('meta_strip'):
        doc.set_metadata({})
        doc.del_xml_metadata()
        return pdf_bytes(doc), f"{base}.pdf", "Todos os metadados foram removidos."
    meta = {key: form.get(field, '').strip() for key, field in META_FIELDS.items()}
    shown = read_metadata(doc)
    dates, moments = {}, {}
    for key, field in META_DATES.items():
        moment = form_date(form.get(field, '').strip())
        if moment and moment == form_date(shown[field]):
            # Não foi mexida: fica como estava, com o fuso que tinha ou sem nenhum (a interface
            # mostra no fuso deste computador, e regravar assim trocaria ou inventaria o fuso)
            dates[key] = doc.metadata[key]
            if parse_pdf_date(dates[key]):  # sem ela no Info, a data veio do XMP, que também fica como está
                moments[key] = parse_pdf_date(dates[key])
        else:
            dates[key], moments[key] = pdf_date(moment) if moment else '', moment
    doc.set_metadata(meta | dates)
    xmp = doc.get_xml_metadata()
    if xmp:
        values = meta | {key: m.isoformat(timespec='seconds') if m else '' for key, m in moments.items()}
        updated = update_xmp(xmp, values)
        if updated:
            doc.set_xml_metadata(updated)
        else:
            doc.del_xml_metadata()  # ilegível: melhor sem XMP do que um contradizendo os campos
    return pdf_bytes(doc), f"{base}.pdf", "Metadados atualizados."


# --- METADADOS DE IMAGEM ---
# A imagem não é regravada (isso recomprimiria o JPEG): só os blocos de metadados do arquivo
# (segmentos do JPEG, chunks do PNG e do WebP) são trocados, e os dados da imagem ficam
# idênticos, byte a byte. EXIF, XMP e textos do PNG são atualizados juntos, porque cada
# programa lê um deles.

IMAGE_FIELDS = ('img_title', 'img_description', 'img_author', 'img_copyright', 'img_keywords', 'img_software', 'img_taken')
IMAGE_EXIF_TEXT = {'img_description': 0x010E, 'img_author': 0x013B, 'img_copyright': 0x8298, 'img_software': 0x0131}
IMAGE_EXIF_XP = {'img_title': 0x9C9B, 'img_author': 0x9C9D, 'img_keywords': 0x9C9E}  # campos do Windows, UTF-16
EXIF_IFD, GPS_IFD, DATE_TAKEN, ORIENTATION, MAKE, MODEL = 0x8769, 0x8825, 0x9003, 0x0112, 0x010F, 0x0110
IMAGE_XMP_PROPERTIES = {
    'img_title': ('dc', 'title', 'Alt'), 'img_description': ('dc', 'description', 'Alt'),
    'img_author': ('dc', 'creator', 'Seq'), 'img_copyright': ('dc', 'rights', 'Alt'),
    'img_keywords': ('dc', 'subject', 'Bag'), 'img_software': ('xmp', 'CreatorTool', 'text'),
    'img_taken': ('exif', 'DateTimeOriginal', 'text'),
}
PNG_TEXT = {'img_title': 'Title', 'img_description': 'Description', 'img_author': 'Author',
            'img_copyright': 'Copyright', 'img_software': 'Software', 'img_taken': 'Creation Time'}
XMP_ID = b'http://ns.adobe.com/xap/1.0/\x00'
# IPTC-IIM, no bloco do Photoshop (APP13) do JPEG: número do dataset no registro 2
IPTC = {'img_title': 5, 'img_description': 120, 'img_author': 80, 'img_copyright': 116, 'img_keywords': 25}
IPTC_DATE, IPTC_TIME, IPTC_UTF8 = 55, 60, b'\x1b%G'
IPTC_BINARY = {0, 125, 200, 201, 202}  # no registro 2, os outros datasets são texto
PHOTOSHOP_ID = b'Photoshop 3.0\x00'


def exif_text(value):
    """Texto de uma tag EXIF: quase todo programa grava UTF-8, que o Pillow lê como Latin-1."""
    value = value.decode('latin-1') if isinstance(value, bytes) else str(value or '')
    try:
        value = value.encode('latin-1').decode('utf-8')
    except UnicodeError:
        pass
    return value.strip('\x00 ')


def exif_xp(value):
    """Campo do Windows (XPTitle etc.): UTF-16."""
    if isinstance(value, tuple):
        value = bytes(value)
    return value.decode('utf-16-le', errors='ignore').rstrip('\x00') if isinstance(value, bytes) else ''


def is_gps(tag):
    return tag.startswith(f"{{{XMP_NS['exif']}}}GPS")


def xmp_values(xmp, properties):
    """Valores atuais das propriedades no XMP; listas viram texto separado por vírgula."""
    try:
        root = ET.fromstring(xmp)
    except ET.ParseError:
        return {}
    rdf = XMP_NS['rdf']
    found = {}
    for field, (prefix, name, _) in properties.items():
        tag = f'{{{XMP_NS[prefix]}}}{name}'
        for description in root.iter(f'{{{rdf}}}Description'):
            if tag in description.attrib:
                found[field] = description.attrib[tag]
                break
            prop = description.find(tag)
            if prop is not None:
                items = list(prop.iter(f'{{{rdf}}}li'))
                if properties[field][2] == 'Alt':  # o mesmo texto em vários idiomas: vale o padrão
                    items = [li for li in items if li.get('{http://www.w3.org/XML/1998/namespace}lang') == 'x-default'] or items[:1]
                items = [li.text or '' for li in items]
                found[field] = ', '.join(items) if items else (prop.text or '').strip()
                break
    return found


def gps_text(gps):
    """Coordenadas em graus decimais, para mostrar."""
    def degrees(value, ref):
        d, m, s = (float(v) for v in value)
        return (-1 if ref in ('S', 'W') else 1) * (d + m / 60 + s / 3600)
    try:
        return f"{degrees(gps[2], gps[1]):.5f}, {degrees(gps[4], gps[3]):.5f}"
    except (KeyError, TypeError, ValueError, ZeroDivisionError):
        return 'sim (sem coordenadas legíveis)'


def photoshop_resources(body):
    """Recursos do bloco do Photoshop como [id, nome, dados]; None se ele não puder ser lido
    inteiro (quando continua em outro segmento, por exemplo)."""
    resources, pos = [], 0
    while body[pos:pos + 4] == b'8BIM' and pos + 7 <= len(body):
        start = pos + 6 + ((body[pos + 6] + 2) & ~1)  # o nome é uma string Pascal de tamanho par
        size = int.from_bytes(body[start:start + 4], 'big')
        if start + 4 + size > len(body):
            return None
        resources.append([body[pos + 4:pos + 6], body[pos + 6:start], body[start + 4:start + 4 + size]])
        pos = start + 4 + size + (size & 1)
    return resources if not body[pos:].strip(b'\x00') else None


def iptc_records(data):
    """Datasets do IPTC como ((registro, dataset), valor, bytes originais)."""
    records, pos = [], 0
    while pos + 5 <= len(data) and data[pos] == 0x1C:
        head, size = 5, int.from_bytes(data[pos + 3:pos + 5], 'big')
        if size & 0x8000:  # tamanho estendido: os bytes seguintes dizem o tamanho
            head += size & 0x7FFF
            size = int.from_bytes(data[pos + 5:pos + head], 'big')
        records.append(((data[pos + 1], data[pos + 2]), data[pos + head:pos + head + size], data[pos:pos + head + size]))
        pos += head + size
    return records


def iptc_values(data):
    """Campos do formulário presentes no IPTC; repetidos (autores, palavras-chave) viram texto separado por vírgula."""
    found = {}
    for (record, dataset), value, _ in iptc_records(data):
        if record == 2:
            found.setdefault(dataset, []).append(exif_text(value))
    values = {field: ', '.join(found[dataset]) for field, dataset in IPTC.items() if dataset in found}
    date, time = found.get(IPTC_DATE, [''])[0], found.get(IPTC_TIME, ['000000'])[0]  # '20260929', '143000-0300'
    if re.fullmatch(r'\d{8}', date) and re.match(r'\d{6}', time):
        values['img_taken'] = f"{date[:4]}-{date[4:6]}-{date[6:]}T{time[:2]}:{time[2:4]}:{time[4:6]}"
    return values


def iptc_dataset(record, dataset, value):
    if len(value) > 0x7FFF:
        raise UserError("Os metadados ficaram grandes demais para um JPEG.")
    return bytes((0x1C, record, dataset)) + len(value).to_bytes(2, 'big') + value


def new_iptc(data, values):
    """IPTC com os campos do formulário no lugar dos antigos, em UTF-8. O resto (cidade, crédito...)
    fica, passado para UTF-8 se estava em outra codificação."""
    records = iptc_records(data)
    utf8 = any(key == (1, 90) and value == IPTC_UTF8 for key, value, _ in records)
    managed = set(IPTC.values()) | {IPTC_DATE, IPTC_TIME}
    out = [((1, 90), iptc_dataset(1, 90, IPTC_UTF8))]
    for (record, dataset), value, raw in records:
        if (record, dataset) == (1, 90) or (record == 2 and dataset in managed):
            continue
        if record == 2 and dataset not in IPTC_BINARY and not utf8:
            raw = iptc_dataset(2, dataset, exif_text(value).encode('utf-8'))
        out.append(((record, dataset), raw))
    for field, dataset in IPTC.items():
        items = values[field].split(',') if field == 'img_keywords' else [values[field]]
        out += [((2, dataset), iptc_dataset(2, dataset, item.strip().encode('utf-8'))) for item in items if item.strip()]
    if values['img_taken']:  # '2026-09-29T14:30:00' -> '20260929' e '143000'
        date, time = values['img_taken'].replace('-', '').replace(':', '').split('T')
        out += [((2, IPTC_DATE), iptc_dataset(2, IPTC_DATE, date.encode())),
                ((2, IPTC_TIME), iptc_dataset(2, IPTC_TIME, time.encode()))]
    return b''.join(raw for _, raw in sorted(out, key=lambda item: item[0]))  # em ordem de registro e dataset


def new_photoshop(body, values):
    """Bloco do Photoshop com o IPTC atualizado e o resto (miniatura, traçados...) igual. Sem IPTC,
    ou se o bloco não puder ser lido inteiro, fica como está."""
    resources = photoshop_resources(body)
    iptc = next((r for r in resources or [] if r[0] == b'\x04\x04'), None)
    if not iptc:
        return body
    iptc[2] = new_iptc(iptc[2], values)
    for resource in resources:
        if resource[0] == b'\x04\x25':  # resumo MD5 do IPTC: sem ele em dia, o IPTC pareceria editado por fora do XMP
            resource[2] = hashlib.md5(iptc[2]).digest()
    return b''.join(b'8BIM' + rid + name + len(d).to_bytes(4, 'big') + d + b'\x00' * (len(d) & 1) for rid, name, d in resources)


def image_parts(data):
    """(formato, EXIF, XMP em texto, textos do PNG, IPTC, tamanho, se tem transparência) sem decodificar a imagem."""
    try:
        img = Image.open(io.BytesIO(data))
        with img:
            fmt = 'JPEG' if img.format == 'MPO' else img.format
            exif = img.getexif()
            xmp = img.info.get('xmp') or img.info.get('XML:com.adobe.xmp') or ''
            if not xmp and fmt == 'JPEG':  # o Pillow antes do 11 não expõe o XMP do JPEG em info
                xmp = next((data[len(XMP_ID):] for marker, data in getattr(img, 'applist', [])
                            if marker == 'APP1' and data.startswith(XMP_ID)), '')
            texts = dict(img.text) if fmt == 'PNG' else {}
            photoshop = next((photoshop_resources(data[len(PHOTOSHOP_ID):]) for marker, data in getattr(img, 'applist', [])
                              if marker == 'APP13' and data.startswith(PHOTOSHOP_ID)), None)
            iptc = next((r[2] for r in photoshop or [] if r[0] == b'\x04\x04'), b'')
            size, alpha = img.size, has_alpha(img)
    except Image.DecompressionBombError:
        raise UserError(TOO_BIG)
    except Exception:
        raise UserError("O arquivo não é uma imagem válida.")
    if fmt not in ('JPEG', 'PNG', 'WEBP'):
        raise UserError("Metadados de imagem: use JPG, PNG ou WebP.")
    return fmt, exif, xmp.decode('utf-8', 'replace') if isinstance(xmp, bytes) else xmp, texts, iptc, size, alpha


MONTHS = {name: i % 12 + 1 for i, name in enumerate(
    'jan fev mar abr mai jun jul ago set out nov dez jan feb mar apr may jun jul aug sep oct nov dec'.split())}


def browser_date(text):
    """Data para o campo do navegador ('2026-09-29T14:30:00'), vinda do EXIF ou ISO ('2026:09:29 14:30:00')
    ou de texto livre, como o 'Creation Time' do PNG: RFC 1123 ('Tue, 29 Sep 2026 14:30:00 GMT'),
    asctime ('Tue Sep 29 14:30:00 2026') ou o idioma do sistema, como grava o gnome-screenshot
    ('ter 29 set 2026 14:30:00', 'Tue 29 Sep 2026 02:30:00 PM'). '' se ilegível."""
    m = re.match(r'(\d{4})[:-](\d\d)[:-](\d\d)[T ](\d\d):(\d\d)(?::(\d\d))?', text)
    if m:
        return f"{m[1]}-{m[2]}-{m[3]}T{m[4]}:{m[5]}:{m[6] or '00'}"
    # por extenso: tirando a hora e o ano, sobram o dia e o nome do mês, em qualquer ordem
    time = re.search(r'(\d{1,2}):(\d\d)(?::(\d\d))?(?:\s*([AaPp])\.?\s?[Mm]\b)?', text)
    year = re.search(r'\b\d{4}\b', text)
    if not (time and year):
        return ''
    rest = text.replace(time[0], ' ').replace(year[0], ' ', 1)
    day = re.search(r'\b\d{1,2}\b', rest)
    month = next((MONTHS[w[:3].lower()] for w in re.findall(r'[^\W\d_]{3,}', rest) if w[:3].lower() in MONTHS), None)
    hour = int(time[1]) % 12 + (12 if time[4] in 'Pp' else 0) if time[4] else int(time[1])
    try:
        return datetime(int(year[0]), month, int(day[0]), hour, int(time[2]), int(time[3] or 0)).isoformat()
    except (TypeError, ValueError):  # sem dia ou mês, ou uma data que não existe
        return ''


def text_date(text):
    """Data escrita fora do padrão do PDF: ISO, como no XMP (o fuso, se houver, fica), ou por extenso."""
    try:
        return datetime.fromisoformat(text.strip().replace('Z', '+00:00'))
    except ValueError:
        iso = browser_date(text)
        return datetime.fromisoformat(iso) if iso else None


def read_image_metadata(data):
    fmt, exif, xmp, texts, iptc, _, _ = image_parts(data)
    values = dict.fromkeys(IMAGE_FIELDS, '') | iptc_values(iptc)  # o EXIF e o XMP, se tiverem o campo, valem mais
    for field, key in PNG_TEXT.items():
        values[field] = texts.get(key, '') or values[field]
    for field, tag in IMAGE_EXIF_TEXT.items():
        values[field] = exif_text(exif.get(tag)) or values[field]
    for field, tag in IMAGE_EXIF_XP.items():
        values[field] = exif_xp(exif.get(tag)) or values[field]
    taken = exif.get_ifd(EXIF_IFD).get(DATE_TAKEN)
    if taken:
        values['img_taken'] = str(taken)
    if xmp:  # o XMP é UTF-8 de verdade: quando tem o campo, vale ele
        values.update({k: v for k, v in xmp_values(xmp, IMAGE_XMP_PROPERTIES).items() if v})
    values['img_taken'] = browser_date(values['img_taken'])
    camera = ' '.join(filter(None, (exif_text(exif.get(MAKE)), exif_text(exif.get(MODEL)))))
    gps = exif.get_ifd(GPS_IFD)
    return values | {'tipo': 'imagem', 'camera': camera, 'localizacao': gps_text(gps) if gps else ''}


def new_image_metadata(exif, xmp, form):
    """(EXIF, XMP, valores do formulário já normalizados) novos a partir do formulário."""
    if form.get('meta_strip'):  # fica só o que muda a aparência: a orientação (o perfil de cor fica no arquivo)
        kept = Image.Exif()
        if ORIENTATION in exif:
            kept[ORIENTATION] = exif[ORIENTATION]
        return kept, None, {}
    values = {field: form.get(field, '').strip() for field in IMAGE_FIELDS}
    taken_iso = ''
    # get_ifd só no bloco Exif que já existe: sem ele, o Pillow 10 não grava o bloco que o
    # get_ifd devolve (a data se perdia) e os mais novos o gravam mesmo vazio
    if values['img_taken']:
        try:
            taken = datetime.fromisoformat(values['img_taken'])
        except ValueError:
            raise UserError(f"Data inválida: \"{values['img_taken']}\".")
        taken_iso = taken.strftime('%Y-%m-%dT%H:%M:%S')
        date = taken.strftime('%Y:%m:%d %H:%M:%S')
        if EXIF_IFD in exif:
            exif.get_ifd(EXIF_IFD)[DATE_TAKEN] = date
        else:
            exif[EXIF_IFD] = {DATE_TAKEN: date}
    elif EXIF_IFD in exif:
        ifd = exif.get_ifd(EXIF_IFD)
        ifd.pop(DATE_TAKEN, None)
        if not ifd:
            del exif[EXIF_IFD]
    for field, tag in IMAGE_EXIF_TEXT.items():
        if values[field]:
            exif[tag] = values[field].encode('utf-8')
        elif tag in exif:
            del exif[tag]
    for field, tag in IMAGE_EXIF_XP.items():
        if values[field]:
            exif[tag] = values[field].encode('utf-16-le') + b'\x00\x00'
        elif tag in exif:
            del exif[tag]
    strip_gps = bool(form.get('img_strip_gps'))
    if strip_gps and GPS_IFD in exif:
        del exif[GPS_IFD]
    values['img_taken'] = taken_iso
    new_xmp = update_xmp(xmp, values, IMAGE_XMP_PROPERTIES, drop=is_gps if strip_gps else None) if xmp else None
    return exif, new_xmp, values


def jpeg_segment(marker, payload):
    if len(payload) > 65533:
        raise UserError("Os metadados ficaram grandes demais para um JPEG.")
    return bytes((0xFF, marker)) + (len(payload) + 2).to_bytes(2, 'big') + payload


def jpeg_image_end(data, pos):
    """Onde termina a imagem principal (logo depois do EOI), a partir do SOS em pos. Nos dados da
    imagem, 0xFF só é marcador se não vier seguido de 0x00 ou de um RST; os segmentos entre as
    varreduras de um JPEG progressivo têm tamanho e são pulados inteiros. Sem EOI, vai até o fim."""
    while pos + 4 <= len(data):
        pos += 2 + int.from_bytes(data[pos + 2:pos + 4], 'big')  # o cabeçalho do segmento
        while (pos := data.find(b'\xff', pos)) >= 0 and pos + 1 < len(data) and \
                (data[pos + 1] in (0x00, 0xFF) or 0xD0 <= data[pos + 1] <= 0xD7):
            pos += 1
        if pos < 0 or pos + 1 >= len(data):
            break
        if data[pos + 1] == 0xD9:
            return pos + 2
    return len(data)


def rewrite_jpeg(data, tiff, xmp, values, strip, drop_extra=False):
    """Troca os segmentos de EXIF e XMP e atualiza o IPTC; os dados da imagem (a partir do SOS) ficam
    iguais. Com drop_extra, sai também o que vem depois da imagem: o vídeo das fotos em movimento, a
    segunda imagem do MPO e o trailer da Samsung guardam os dados deles, localização inclusive.
    Devolve (JPEG, se havia algo depois)."""
    head, others, pos = [], [], 2
    while True:
        if pos + 4 > len(data) or data[pos] != 0xFF:
            raise UserError("O JPEG está corrompido.")
        marker = data[pos + 1]
        if marker == 0xDA:  # início dos dados da imagem: daqui em diante nada muda
            rest = data[pos:]
            break
        end = pos + 2 + int.from_bytes(data[pos + 2:pos + 4], 'big')
        segment, body = data[pos:end], data[pos + 4:end]
        pos = end
        if marker == 0xE1 and (body.startswith(b'Exif\x00\x00') or body.startswith(XMP_ID)
                               or body.startswith(b'http://ns.adobe.com/xmp/extension/')):
            continue  # EXIF e XMP antigos: saem, os novos entram abaixo
        if strip and (marker == 0xED or marker == 0xFE):  # IPTC (Photoshop) e comentário
            continue
        if drop_extra and marker == 0xE2 and body.startswith(b'MPF\x00'):  # o índice das imagens que saem
            continue
        if marker == 0xED and body.startswith(PHOTOSHOP_ID):
            segment = jpeg_segment(0xED, PHOTOSHOP_ID + new_photoshop(body[len(PHOTOSHOP_ID):], values))
        (head if marker == 0xE0 and not others else others).append(segment)  # JFIF fica primeiro
    new = []
    if tiff:
        new.append(jpeg_segment(0xE1, b'Exif\x00\x00' + tiff))
    if xmp:
        new.append(jpeg_segment(0xE1, XMP_ID + xmp.encode('utf-8')))
    extra = b''
    if drop_extra:
        end = jpeg_image_end(rest, 0)
        rest, extra = rest[:end], rest[end:]
    return b'\xff\xd8' + b''.join(head + new + others) + rest, bool(extra.strip(b'\x00'))


def png_chunk(kind, payload):
    return len(payload).to_bytes(4, 'big') + kind + payload + zlib.crc32(kind + payload).to_bytes(4, 'big')


def rewrite_png(data, tiff, xmp, values, strip):
    """Troca o eXIf e os textos de metadados; IDAT e o resto ficam iguais."""
    texts = {key: values[field] for field, key in PNG_TEXT.items() if values.get(field)}
    managed = {key.encode('latin-1') for key in PNG_TEXT.values()} | {b'XML:com.adobe.xmp'}
    new = []
    if tiff:
        new.append(png_chunk(b'eXIf', tiff))
    for key, text in list(texts.items()) + ([('XML:com.adobe.xmp', xmp)] if xmp else []):
        # iTXt: palavra-chave, sem compressão, idioma e tradução vazios, texto em UTF-8
        new.append(png_chunk(b'iTXt', key.encode('latin-1') + b'\x00\x00\x00\x00\x00' + text.encode('utf-8')))
    out, pos = [data[:8]], 8
    while pos < len(data):
        length = int.from_bytes(data[pos:pos + 4], 'big')
        kind, chunk = data[pos + 4:pos + 8], data[pos:pos + 12 + length]
        pos += 12 + length
        if kind == b'eXIf' or (strip and kind == b'tIME'):
            continue
        if kind in (b'tEXt', b'zTXt', b'iTXt') and (strip or chunk[8:].split(b'\x00', 1)[0] in managed):
            continue
        if kind == b'IDAT' and new:  # metadados antes dos dados da imagem
            out.extend(new)
            new = []
        out.append(chunk)
    return b''.join(out)


def rewrite_webp(data, tiff, xmp, size, alpha):
    """Troca os chunks EXIF e XMP; os dados da imagem ficam iguais. Um WebP simples ganha o
    cabeçalho estendido (VP8X), que é onde se declara que há metadados."""
    chunks, pos = [], 12
    while pos + 8 <= len(data):
        kind, length = data[pos:pos + 4], int.from_bytes(data[pos + 4:pos + 8], 'little')
        chunks.append((kind, data[pos + 8:pos + 8 + length]))
        pos += 8 + length + (length & 1)
    chunks = [(k, c) for k, c in chunks if k not in (b'EXIF', b'XMP ')]
    if tiff:
        chunks.append((b'EXIF', tiff))
    if xmp:
        chunks.append((b'XMP ', xmp.encode('utf-8')))
    if chunks and chunks[0][0] == b'VP8X':
        flags = chunks[0][1][0] & ~0x0C
    elif tiff or xmp:
        flags = 0x10 if alpha else 0  # WebP simples: cria o VP8X
        chunks.insert(0, (b'VP8X', bytes(4) + (size[0] - 1).to_bytes(3, 'little') + (size[1] - 1).to_bytes(3, 'little')))
    else:
        flags = None
    if flags is not None:
        flags |= (0x08 if tiff else 0) | (0x04 if xmp else 0)
        chunks[0] = (b'VP8X', bytes([flags]) + chunks[0][1][1:])
    body = b'WEBP' + b''.join(k + len(c).to_bytes(4, 'little') + c + (b'\x00' if len(c) & 1 else b'') for k, c in chunks)
    return b'RIFF' + len(body).to_bytes(4, 'little') + body


def edit_image_metadata(path, base, form):
    data = path.read_bytes()
    fmt, exif, xmp, _, _, size, alpha = image_parts(data)
    exif, new_xmp, values = new_image_metadata(exif, xmp, form)
    tiff = exif.tobytes()[6:] if len(exif) else None  # sem o prefixo "Exif\0\0" do JPEG
    strip = bool(form.get('meta_strip'))
    extra = False
    if fmt == 'JPEG':
        out, extra = rewrite_jpeg(data, tiff, new_xmp, values, strip, drop_extra=strip or bool(form.get('img_strip_gps')))
    elif fmt == 'PNG':
        out = rewrite_png(data, tiff, new_xmp, values, strip)
    else:
        out = rewrite_webp(data, tiff, new_xmp, size, alpha)
    message = "Todos os metadados foram removidos (a orientação e o perfil de cor ficaram)." if strip else "Metadados atualizados."
    if extra:
        message += " O que vinha grudado depois da foto (vídeo da foto em movimento, imagens extras) também saiu."
    return out, f"{base}{path.suffix}", message


def number_pages(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path)
    label = NUMBER_FORMATS.get(form.get('format'), NUMBER_FORMATS['n'])
    for page in doc:
        text = label.format(n=page.number + 1, t=doc.page_count)
        r = page.rect  # área visível, já considerando a rotação da página
        x = (r.width - pymupdf.get_text_length(text, 'helv', 10)) / 2
        # Centralizado no rodapé como o usuário vê, mesmo em páginas giradas
        page.insert_text(pymupdf.Point(x, r.height - 20) * page.derotation_matrix, text,
                         fontsize=10, fontname='helv', rotate=page.rotation)
    return pdf_bytes(doc), f"{base}_numerado.pdf"


def add_invisible_text(page, words):
    """Escreve as palavras do OCR invisíveis sobre a imagem, para buscar e selecionar."""
    # Todas as palavras de uma linha usam a mesma base e altura, para a linha não se partir
    lines = {}
    for x0, y0, x1, y1, word, block, line, _ in words:
        bottom, height = lines.get((block, line), (0, 0))
        lines[(block, line)] = (max(bottom, y1), max(height, y1 - y0))
    for x0, y0, x1, y1, word, block, line, _ in words:
        bottom, height = lines[(block, line)]
        word = word.encode('cp1252', 'replace').decode('cp1252')  # a fonte padrão só tem esses caracteres
        if x1 <= x0 or height <= 0 or not word.strip():
            continue  # caixa degenerada do OCR: não há onde escrever
        start = pymupdf.Point(x0, bottom - 0.2 * height)
        stretch = (x1 - x0) / pymupdf.get_text_length(word, 'helv', height)  # ocupa a largura da palavra
        page.insert_text(start, word + ' ', fontsize=height, fontname='helv', render_mode=3,
                         morph=(start, pymupdf.Matrix(stretch, 1)))


def ocr_pdf(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path)
    tessdata = find_tessdata()
    languages = [lang for lang in ('por', 'eng') if tessdata and Path(tessdata, f'{lang}.traineddata').exists()]
    if 'por' not in languages:
        raise UserError("O OCR precisa do Tesseract instalado com o idioma português. Veja as instruções no README.")
    recognized = 0
    for page in doc:
        if page.get_text().strip():
            continue  # a página já tem texto
        page.remove_rotation()  # mesma aparência, mas com coordenadas iguais às que o OCR devolve
        textpage = page.get_textpage_ocr(dpi=300, language='+'.join(languages), tessdata=tessdata, full=True)
        add_invisible_text(page, page.get_text('words', textpage=textpage))
        recognized += 1
    if not recognized:
        raise UserError("Todas as páginas deste PDF já têm texto: não há o que reconhecer.")
    return pdf_bytes(doc), f"{base}_ocr.pdf", f"Texto reconhecido em {recognized} página(s)."


def format_size(size):
    return f"{size / 1024:.0f} KB" if size < 1024 * 1024 else f"{size / 1024 / 1024:.1f} MB".replace('.', ',')


PDF_LEVELS = {'recomendada': (150, 75), 'forte': (96, 55)}  # (DPI alvo, qualidade do JPEG)


def pdf_images(doc):
    """Imagens que a compressão pode mexer: {xref: (página, largura, altura, menor DPI em que aparecem)}."""
    found = {}
    for page in doc:
        for item in page.get_images(full=True):
            xref, smask, width, height, bpc = item[:5]
            if smask or bpc == 1 or doc.xref_get_key(xref, 'Mask')[0] != 'null':
                continue  # transparência e imagens de 1 bit (texto escaneado em preto e branco) ficam como estão
            rect = page.get_image_bbox(item)  # bem mais rápido que get_image_rects, que decodifica a imagem
            if rect.is_empty or rect.is_infinite:
                continue
            dpi = 72 * math.sqrt(width * height / (rect.width * rect.height))  # vale também para imagem girada
            if xref not in found or dpi < found[xref][3]:
                found[xref] = (page.number, width, height, dpi)
    return found


def stored_size(doc, xref, raw=None):
    """Tamanho da imagem como fica no PDF gravado: stream sem compressão ganha deflate ao salvar."""
    raw = doc.xref_stream_raw(xref) if raw is None else raw
    return len(zlib.compress(raw)) if doc.xref_get_key(xref, 'Filter')[0] == 'null' else len(raw)


def shrink_pdf_image(doc, xref, width, height, dpi, target_dpi, quality):
    """A imagem em JPEG, reduzida à resolução alvo; None se isso não a deixar menor."""
    raw = doc.xref_stream_raw(xref)
    is_jpeg = doc.xref_get_key(xref, 'Filter')[1] == '/DCTDecode'
    if dpi <= target_dpi + 10 and not is_jpeg:
        return None  # imagem sem perda (desenho, captura de tela) só vira JPEG se tiver resolução demais
    scale = min(1, target_dpi / dpi)
    size = (max(1, round(width * scale)), max(1, round(height * scale)))
    try:
        img = None
        if is_jpeg and doc.xref_get_key(xref, 'Decode')[0] == 'null':
            img = Image.open(io.BytesIO(raw))
            if img.mode in ('RGB', 'L'):
                img.draft(img.mode, size)  # decodifica o JPEG já reduzido: bem mais rápido
            else:
                img = None  # CMYK e afins: o MuPDF converte as cores melhor
        if img is None:
            pix = pymupdf.Pixmap(doc, xref)
            if pix.n not in (1, 3):
                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
            img = pix.pil_image()
        img = img.convert('L' if img.mode == 'L' else 'RGB')
    except Exception:
        return None  # formato de imagem que não dá para decodificar: fica como está
    if img.size != size:
        img = img.resize(size, Image.Resampling.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, 'JPEG', quality=quality, optimize=True)
    return buf.getvalue() if buf.tell() < stored_size(doc, xref, raw) else None


def compress_pdf(files, form, tmp):
    path, base = files[0]
    doc = open_pdf(path)
    target_dpi, quality = PDF_LEVELS.get(form.get('level'), PDF_LEVELS['recomendada'])
    for xref, (page_number, width, height, dpi) in pdf_images(doc).items():
        smaller = shrink_pdf_image(doc, xref, width, height, dpi, target_dpi, quality)
        if smaller:
            doc[page_number].replace_image(xref, stream=smaller)  # vale para todas as páginas que a usam
    data = doc.tobytes(garbage=4, deflate=True, clean=True, use_objstms=1)  # garbage=4 junta as cópias
    before = path.stat().st_size
    if len(data) >= before:
        return path.read_bytes(), f"{base}.pdf", "Este PDF já está otimizado: não deu para reduzir mais."
    return data, f"{base}_comprimido.pdf", \
        f"Reduzido de {format_size(before)} para {format_size(len(data))} ({1 - len(data) / before:.0%} menor)."


# nível: (qualidade do JPG e do WebP, cores do PNG; None = PNG sem perda)
IMAGE_LEVELS = {
    'leve': (85, None),
    'recomendada': (72, 256),
    'forte': (55, 128),
    'extrema': (35, 64),
}
IMAGE_EXTENSIONS = {'JPEG': '.jpg', 'PNG': '.png', 'WEBP': '.webp'}
MAX_SIDES = {'2560': 2560, '1920': 1920, '1280': 1280, '800': 800}


def image_settings(form):
    """Opções do Comprimir imagem: (qualidade, cores do PNG, lado máximo, formato, remover dados da foto)."""
    level = form.get('image_level') if form.get('image_level') in IMAGE_LEVELS else 'recomendada'
    return (*IMAGE_LEVELS[level], MAX_SIDES.get(form.get('max_size')),
            form.get('image_format', '').upper(), bool(form.get('strip_metadata')))


def prepare_image(path, max_side, target, strip_metadata):
    """A imagem já reduzida e o que é preciso para gravá-la:
    (imagem, formato final, opções de gravação, se muda algo além da compressão)."""
    img, fmt = open_image(path, reduce_to=max_side)
    out_fmt = target if target in IMAGE_EXTENSIONS else fmt if fmt in IMAGE_EXTENSIONS else 'PNG'
    resized = bool(max_side) and max(img.size) > max_side
    if resized:
        img.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)

    options = {}
    if img.info.get('icc_profile'):
        options['icc_profile'] = img.info['icc_profile']  # perfil de cor: sem ele as cores mudam
    has_exif = bool(img.getexif())
    if has_exif and not strip_metadata:
        options['exif'] = img.getexif()  # a orientação já foi aplicada e retirada
    return img, out_fmt, options, out_fmt != fmt or resized or (has_exif and strip_metadata)


def encode_image(img, fmt, quality, colors, options=None, fast_png=False):
    options = options or {}
    buf = io.BytesIO()
    if fmt == 'JPEG':
        on_white(img).save(buf, 'JPEG', quality=quality, optimize=True, progressive=True, **options)
    elif fmt == 'WEBP':
        # method 6 é o mais compacto, mas com transparência fica ~10 vezes mais lento para ganhar 4 a 5%
        img = img.convert('RGBA' if has_alpha(img) else 'RGB')
        img.save(buf, 'WEBP', quality=quality, method=5 if img.mode == 'RGBA' else 6, **options)
    else:
        img = img.convert('RGBA' if has_alpha(img) else 'RGB')
        if colors:  # PNG com menos cores, como faz o TinyPNG
            method = Image.Quantize.FASTOCTREE if img.mode == 'RGBA' else Image.Quantize.MEDIANCUT
            img = img.quantize(colors, method=method, dither=Image.Dither.FLOYDSTEINBERG)
        img.save(buf, 'PNG', **({'compress_level': 6} if fast_png else {'optimize': True}), **options)
    return buf.getvalue()


def compress_image(path, settings):
    """Devolve (bytes, extensão, se ficou como estava)."""
    quality, colors, max_side, target, strip_metadata = settings
    img, fmt, options, changed = prepare_image(path, max_side, target, strip_metadata)
    data = encode_image(img, fmt, quality, colors, options)
    # Se não deu para diminuir e não havia nada a mudar (formato, tamanho, dados da foto), fica o original
    if len(data) >= path.stat().st_size and not changed:
        return path.read_bytes(), IMAGE_EXTENSIONS[fmt], True
    return data, IMAGE_EXTENSIONS[fmt], False


def pack(results, zip_name):
    """Um resultado vai direto; vários vão num ZIP. results: [(nome sem extensão, extensão, bytes)]."""
    if len(results) == 1:
        stem, ext, data = results[0]
        return data, f"{stem}{ext}"
    buf, used = io.BytesIO(), set()
    with zipfile.ZipFile(buf, 'w') as z:  # imagens já vêm comprimidas
        for stem, ext, data in results:
            name, n = f"{stem}{ext}", 2
            while name in used:  # dois arquivos com o mesmo nome não podem se sobrescrever no ZIP
                name, n = f"{stem}_{n}{ext}", n + 1
            used.add(name)
            z.writestr(name, data)
    return buf.getvalue(), zip_name


def compress_images(files, form, tmp):
    settings = image_settings(form)
    results, kept = [], 0
    before = after = 0
    for path, base in files:
        data, ext, unchanged = compress_image(path, settings)
        results.append((f"{base}_comprimida", ext, data))
        kept += unchanged
        before += path.stat().st_size
        after += len(data)

    if len(files) == 1 and kept:
        message = "A imagem já estava bem comprimida e ficou como estava."
    else:
        # Trocar o formato (uma foto JPG em PNG sem perda, por exemplo) pode deixar o arquivo maior
        grew = after > before
        change = f"{after / before - 1:.0%} maior" if grew else f"{1 - after / before:.0%} menor"
        prefix = f"{len(files)} imagens: de" if len(files) > 1 else "Ficou maior: de" if grew else "Reduzida de"
        message = f"{prefix} {format_size(before)} para {format_size(after)} ({change})."
        if kept == 1:
            message += " Uma já estava bem comprimida e ficou como estava."
        elif kept:
            message += f" {kept} já estavam bem comprimidas e ficaram como estavam."
    return *pack(results, "Imagens_comprimidas.zip"), message


# --- REMOVER FUNDO ---
# Modelo de IA ISNet (projeto DIS, licença Apache-2.0), no formato ONNX publicado pelo rembg.
# É baixado só na primeira vez e conferido pelo SHA-256; as imagens nunca saem do computador.
# Dos modelos testados, foi o de melhor recorte que roda bem num computador comum: ~1 s por
# foto e ~800 MB de memória (o BiRefNet passou de 8 GB).
BG_MODEL_URL = 'https://github.com/danielgatis/rembg/releases/download/v0.0.0/isnet-general-use.onnx'
BG_MODEL_SHA256 = '60920e99c45464f2ba57bee2ad08c919a52bbf852739e96947fbb4358c0d964a'
BG_MODEL_MB = 170
BG_MODEL_SIDE = 1024  # o modelo vê a imagem em 1024x1024; a máscara volta ao tamanho original
_bg_session = None
_bg_lock = threading.Lock()


def bg_model_path():
    """O embutido no executável, se houver; senão na pasta de cache do sistema: %LOCALAPPDATA% no Windows, ~/.cache no Linux."""
    if path := bundled('isnet-general-use.onnx'):
        return path
    base = os.environ.get('LOCALAPPDATA') or os.environ.get('XDG_CACHE_HOME') or Path.home() / '.cache'
    return Path(base) / 'euamopdf' / 'isnet-general-use.onnx'


def download_bg_model(path):
    """Baixa o modelo para um arquivo temporário e só o põe no lugar se o SHA-256 conferir."""
    path.parent.mkdir(parents=True, exist_ok=True)
    partial = path.with_name(path.name + '.part')
    digest = hashlib.sha256()
    try:
        with urllib.request.urlopen(BG_MODEL_URL, timeout=30) as response, open(partial, 'wb') as out:
            while chunk := response.read(1 << 20):
                digest.update(chunk)
                out.write(chunk)
    except OSError:
        partial.unlink(missing_ok=True)
        raise UserError("Não foi possível baixar o modelo de IA que remove o fundo. "
                        "Só na primeira vez é preciso estar conectado à internet.")
    if digest.hexdigest() != BG_MODEL_SHA256:
        partial.unlink(missing_ok=True)
        raise UserError("O modelo de IA baixado veio corrompido. Tente de novo.")
    os.replace(partial, path)


def bg_session():
    """Sessão do modelo: baixado na primeira vez e carregado uma vez só."""
    global _bg_session
    with _bg_lock:  # também impede dois downloads ao mesmo tempo
        if _bg_session is None:
            import onnxruntime  # importação lenta: só quando usada
            path = bg_model_path()
            if not path.exists():
                download_bg_model(path)
            options = onnxruntime.SessionOptions()
            options.enable_cpu_mem_arena = False  # pico de ~800 MB em vez de ~1,2 GB, por ~0,3 s a mais
            _bg_session = onnxruntime.InferenceSession(str(path), options, providers=['CPUExecutionProvider'])
        return _bg_session


def subject_mask(img):
    """Máscara do objeto principal (0 a 255), no tamanho da imagem."""
    session = bg_session()
    small = img.convert('RGB').resize((BG_MODEL_SIDE, BG_MODEL_SIDE), Image.Resampling.LANCZOS)
    x = np.asarray(small, dtype=np.float32) / 255
    x = (x - 0.5).transpose(2, 0, 1)[None]  # normalização do ISNet: média 0,5 e desvio 1, canais primeiro
    out = session.run(None, {session.get_inputs()[0].name: x})[0][0, 0]
    # A confiança do modelo fica em ~0,985 no objeto e raramente é 0 no fundo: esticar 0,05-0,95
    # para 0-1 deixa o objeto opaco e o fundo limpo, mantendo a borda suave
    out = (out - 0.05) / 0.9
    mask = Image.fromarray((np.clip(out, 0, 1) * 255).round().astype(np.uint8))
    # Ampliação bicúbica: nos testes, refinar a borda com guided filter piorou o recorte
    return mask.resize(img.size, Image.Resampling.BICUBIC)


BACKGROUNDS = {'branco': '#ffffff', 'preto': '#000000'}
BG_FORMATS = {'png': ('PNG', '.png', {}), 'webp': ('WEBP', '.webp', {'lossless': True})}  # sempre sem perda


def background_color(form):
    """Cor do novo fundo, ou None para transparente."""
    choice = form.get('background', 'transparente')
    if choice != 'cor':
        return BACKGROUNDS.get(choice)
    color = form.get('background_color', '')
    if not re.fullmatch(r'#[0-9a-fA-F]{6}', color):
        raise UserError("Escolha uma cor válida para o fundo.")
    return color


def load_mask(upload, size, what="máscara"):
    """Máscara em tons de cinza enviada pela interface; tem de ter o tamanho da imagem."""
    if upload is None:
        raise UserError(f"Falta a {what}.")
    try:
        mask = Image.open(upload.stream)
        mask.load()
    except Exception:
        raise UserError(f"A {what} enviada não é uma imagem válida.")
    if mask.size != size:
        raise UserError(f"A {what} não tem o tamanho da imagem.")
    return mask.convert('L')


def load_strokes(upload, size, exact=False):
    """Traço pintado na prévia (onde o canal alfa não é zero), levado ao tamanho da imagem: sim
    ou não por pixel, para a detecção; com exact, de 0 a 255, com a borda suavizada ao ampliar."""
    if upload is None:
        raise UserError("Falta o traço.")
    try:
        drawn = Image.open(upload.stream).convert('RGBA').getchannel('A').point(lambda v: 255 if v else 0)
    except Exception:
        raise UserError("O traço enviado não é uma imagem válida.")
    if exact:
        painted = np.asarray(drawn.resize(size, Image.Resampling.BILINEAR))
    else:
        painted = np.asarray(drawn.resize(size, Image.Resampling.NEAREST)) > 0
    if not painted.any():
        raise UserError("Pinte por cima da área que quer corrigir.")
    return painted


def paint_mask(mask, painted, restore):
    """Pincel exato, sem detecção: muda só o que foi pintado."""
    m = np.asarray(mask, dtype=np.uint8)
    return Image.fromarray(np.maximum(m, painted) if restore else np.minimum(m, 255 - painted))


REFINE_SIDE = 1024  # o GrabCut roda num recorte de até 1024 px em volta do traço


def grabcut_region(img, mask, painted, restore, box, reach):
    """Região a mudar dentro do recorte box, no tamanho do recorte."""
    import cv2  # importação lenta: só quando usada
    x0, y0, x1, y1 = box
    scale = min(1, REFINE_SIDE / max(x1 - x0, y1 - y0))
    size = (max(1, round((x1 - x0) * scale)), max(1, round((y1 - y0) * scale)))
    crop = np.asarray(img.convert('RGB').crop(box).resize(size, Image.Resampling.BILINEAR))[..., ::-1].copy()
    cm = np.asarray(Image.fromarray(mask[y0:y1, x0:x1]).resize(size, Image.Resampling.BILINEAR))
    # BOX + "> 0": ao reduzir o recorte, um traço fino não some
    stroke = Image.fromarray(painted[y0:y1, x0:x1].astype(np.uint8) * 255)
    cs = np.asarray(stroke.resize(size, Image.Resampling.BOX)) > 0
    T, PT, O = (cv2.GC_FGD, cv2.GC_PR_FGD, cv2.GC_BGD) if restore else (cv2.GC_BGD, cv2.GC_PR_BGD, cv2.GC_FGD)
    # Só o que já está, com certeza, do lado pedido fica de fora: o meio-termo (semitransparente)
    # também pode ser corrigido
    done = cm >= 250 if restore else cm <= 5
    near = cv2.distanceTransform((~cs).astype(np.uint8), cv2.DIST_L2, 3) <= reach * scale
    # O lado pedido ensina ao GrabCut as cores dele; o resto começa do lado em que o modelo o pôs
    # e pode mudar. Longe do traço, o resto fica travado como está.
    labels = np.where(cm > 127, cv2.GC_PR_FGD, cv2.GC_PR_BGD).astype(np.uint8)
    far = ~near & ~done
    labels[far] = np.where(cm[far] > 127, cv2.GC_FGD, cv2.GC_BGD)
    labels[done] = T
    # Só o miolo do traço é certeza: a borda do pincel, que costuma vazar, o GrabCut decide
    depth = cv2.distanceTransform(cs.astype(np.uint8), cv2.DIST_L2, 3)
    sure = cs & (depth >= 0.5 * depth.max())
    labels[cs & ~done] = PT
    labels[sure] = T
    cv2.setRNGSeed(0)  # mesmo traço, mesmo resultado
    try:
        cv2.grabCut(crop, labels, None, np.zeros((1, 65)), np.zeros((1, 65)), 5, cv2.GC_INIT_WITH_MASK)
    except cv2.error:
        pass  # recorte só de um lado: fica só o miolo do traço
    wanted = np.isin(labels, (T, PT))
    # Muda o que o GrabCut pôs do lado pedido, perto do traço e ligado a ele
    to_change = (wanted & ~done & near) | sure
    _, parts = cv2.connectedComponents(to_change.astype(np.uint8), connectivity=8)
    touched = np.unique(parts[sure])
    region = np.isin(parts, touched[touched > 0])
    full = Image.fromarray(region.astype(np.uint8) * 255).resize((x1 - x0, y1 - y0), Image.Resampling.BILINEAR)
    return np.asarray(full)


def refine_mask(img, mask, painted, restore):
    """Estende o traço do usuário à região que ele quis marcar (restaurar ou apagar) com o
    GrabCut, e devolve a máscara corrigida. O resto da máscara não muda.

    Parâmetros escolhidos numa bancada de 220 erros simulados sobre máscaras conhecidas (pedaços
    do objeto apagados, fundo sobrando, trechos semitransparentes, traços que vazam, um clique
    só): 82% corrigidos por inteiro sem estrago em volta, contra 34% da versão anterior. Deixar
    o alcance crescer, ou partir da semelhança com a cor do traço, espalhava a correção pelo
    fundo de cor parecida."""
    m = np.asarray(mask, dtype=np.uint8)
    h, w = m.shape
    ys, xs = np.nonzero(painted)
    extent = int(max(np.ptp(xs), np.ptp(ys))) + 1
    reach = max(int(0.04 * max(h, w)), 3 * extent)  # até onde a correção pode ir a partir do traço
    pad = 2 * reach  # o recorte mostra o dobro do alcance: o GrabCut vê os dois lados em volta
    box = (max(0, xs.min() - pad), max(0, ys.min() - pad), min(w, xs.max() + pad + 1), min(h, ys.max() + pad + 1))
    region = grabcut_region(img, m, painted, restore, box, reach)
    x0, y0, x1, y1 = box
    out = m.copy()
    part = out[y0:y1, x0:x1]
    out[y0:y1, x0:x1] = np.maximum(part, region) if restore else np.minimum(part, 255 - region)
    return Image.fromarray(out)


def remove_background(path, background, fmt, mask_upload=None):
    """A imagem original sem o fundo, na mesma resolução e com os mesmos pixels no objeto:
    (bytes, extensão, se foi encontrado um objeto em destaque). Com mask_upload, usa a máscara
    corrigida na prévia em vez de rodar o modelo."""
    img, _ = open_image(path)
    # Perfil de cor CMYK ou de cinza não serve para a imagem RGBA gravada
    icc = img.info.get('icc_profile') if img.mode in ('RGB', 'RGBA', 'P') else None
    alpha = load_mask(mask_upload, img.size) if mask_upload else subject_mask(img)
    if has_alpha(img):  # a transparência que a imagem já tinha continua valendo
        alpha = ImageChops.multiply(alpha, img.convert('RGBA').getchannel('A'))
    out = img.convert('RGBA')
    out.putalpha(alpha)
    if background:
        canvas = Image.new('RGB', img.size, background)
        canvas.paste(out, mask=alpha)
        out = canvas
    else:
        # Pixel totalmente transparente guardaria a cor do fundo original, que qualquer um veria
        # tirando a transparência. A borda semitransparente mantém a cor da fonte.
        out = Image.composite(out, Image.new('RGBA', img.size), alpha.point(lambda v: 255 if v else 0))
    pil_fmt, ext, options = BG_FORMATS[fmt]
    buf = io.BytesIO()
    out.save(buf, pil_fmt, **options, **({'icc_profile': icc} if icc else {}))
    return buf.getvalue(), ext, alpha.getextrema()[1] >= 128


def remove_backgrounds(files, form, tmp):
    background = background_color(form)
    fmt = form.get('bg_format') if form.get('bg_format') in BG_FORMATS else 'png'
    masks = request.files.getlist('mascara')  # corrigidas na prévia, uma por imagem, na mesma ordem
    if masks and len(masks) != len(files):
        raise UserError("Envie uma máscara para cada imagem.")
    results, missed = [], []
    for i, (path, base) in enumerate(files):
        data, ext, found = remove_background(path, background, fmt, masks[i] if masks else None)
        results.append((f"{base}_sem_fundo", ext, data))
        if not found:
            missed.append(base)
    message = f"Fundo removido de {len(files)} imagem(ns)."
    if missed:
        message = f"Nenhum objeto em destaque encontrado em: {', '.join(missed)}. Confira o resultado."
    return *pack(results, "Imagens_sem_fundo.zip"), message


# --- ESTIMATIVA DE TAMANHO ---
# Rápida o bastante para rodar a cada mudança de opção: comprime de verdade só uma amostra
# e extrapola. Devolve (tamanho atual, tamanho estimado) em bytes.

def spread(items, n):
    """Até n itens espalhados pela lista."""
    return items if len(items) <= n else [items[i * len(items) // n] for i in range(n)]


# O optimize do PNG deixa o arquivo ~5% menor que o nível padrão, mas é 5 a 10 vezes mais lento
PNG_OPTIMIZE_GAIN = 0.95


def mosaic(img, tile=64, grid=23):
    """Pedaços espalhados pela imagem, montados numa amostra de ~2 MP: (amostra, quantas vezes a
    imagem é maior que ela). Imagens pequenas vão inteiras. Com fotos reais, o erro médio da
    estimativa fica em ~5%; pedaços maiores (e em menor número) erram mais."""
    w, h = img.size
    if w * h <= 2 * (tile * grid) ** 2:
        return img, 1
    img = img.convert('RGBA' if has_alpha(img) else 'L' if img.mode == 'L' else 'RGB')
    sample = Image.new(img.mode, (tile * grid, tile * grid))
    for i in range(grid):
        for j in range(grid):
            x, y = (w - tile) * i // (grid - 1), (h - tile) * j // (grid - 1)
            sample.paste(img.crop((x, y, x + tile, y + tile)), (i * tile, j * tile))
    return sample, w * h / (tile * grid) ** 2


def estimate_image(path, settings):
    quality, colors, max_side, target, strip_metadata = settings
    img, fmt, options, changed = prepare_image(path, max_side, target, strip_metadata)
    sample, scale = mosaic(img)
    # Perfil de cor e EXIF não crescem com a imagem: entram uma vez só
    extra = sum(len(value if isinstance(value, bytes) else value.tobytes()) for value in options.values())
    lossless_png = fmt == 'PNG' and not colors
    size = len(encode_image(sample, fmt, quality, colors, fast_png=lossless_png)) * scale + extra
    if lossless_png:
        size *= PNG_OPTIMIZE_GAIN
    before = path.stat().st_size
    return before if size >= before and not changed else size  # como no compress_image


ESTIMATE_SECONDS = 2  # tempo para estimar imagens; as que não couberem são extrapoladas
ESTIMATE_MIN_IMAGES = 3  # mesmo passando do tempo: extrapolar de uma imagem só erra muito num lote misturado


def estimate_images(files, form):
    settings = image_settings(form)
    sample_before = sample_after = 0
    deadline = time.monotonic() + ESTIMATE_SECONDS
    order = random.Random(0).sample(files, len(files))  # embaralhada: a amostra fica espalhada
    for count, (path, _) in enumerate(order, 1):
        sample_before += path.stat().st_size
        sample_after += estimate_image(path, settings)
        if count >= ESTIMATE_MIN_IMAGES and time.monotonic() > deadline:
            break
    before = sum(path.stat().st_size for path, _ in files)
    return before, before * sample_after / sample_before


def estimate_pdf(files, form):
    path, _ = files[0]
    doc = open_pdf(path)
    target_dpi, quality = PDF_LEVELS.get(form.get('level'), PDF_LEVELS['recomendada'])
    images = pdf_images(doc)
    raw = {xref: stored_size(doc, xref) for xref in images}
    sample = spread(list(images.items()), 8)
    sample_before = sum(raw[xref] for xref, _ in sample)
    sample_after = 0
    for xref, (_, width, height, dpi) in sample:
        smaller = shrink_pdf_image(doc, xref, width, height, dpi, target_dpi, quality)
        sample_after += len(smaller) if smaller else raw[xref]
    # O PDF limpo como o compress_pdf grava, trocando o tamanho das imagens pelo estimado
    cleaned = len(doc.tobytes(garbage=4, deflate=True, clean=True, use_objstms=1))
    images_before = sum(raw.values())
    after = cleaned - images_before + images_before * (sample_after / sample_before if sample_before else 1)
    before = path.stat().st_size
    return before, min(before, after)  # o compress_pdf devolve o original quando não diminui


ESTIMATORS = {
    'compress-pdf': estimate_pdf,
    'compress-image': estimate_images,
}


PDF = ('.pdf',)
WORD = ('.doc', '.docx', '.odt', '.rtf')
EXCEL = ('.xls', '.xlsx', '.ods', '.csv')
POWERPOINT = ('.ppt', '.pptx', '.odp')
IMAGES = ('.jpg', '.jpeg', '.png', '.webp', '.bmp', '.gif', '.tif', '.tiff')
WEB_IMAGES = ('.jpg', '.jpeg', '.png', '.webp')

# ação: (função, extensões aceitas, aceita vários arquivos)
ACTIONS = {
    "merge-pdf": (merge_pdf, PDF, True),
    "compress-pdf": (compress_pdf, PDF, False),
    "compress-image": (compress_images, WEB_IMAGES, True),
    "remove-background": (remove_backgrounds, WEB_IMAGES, True),
    "ocr-pdf": (ocr_pdf, PDF, False),
    "protect-pdf": (protect_pdf, PDF, False),
    "unlock-pdf": (unlock_pdf, PDF, False),
    "watermark-pdf": (watermark_pdf, PDF, False),
    "number-pages": (number_pages, PDF, False),
    "edit-metadata": (edit_metadata, PDF + WEB_IMAGES, False),
    "split-pdf": (split_pdf, PDF, False),
    "rotate-pdf": (rotate_pdf, PDF, False),
    "organize-pdf": (organize_pdf, PDF, False),
    "word-to-pdf": (office_action("Word.Application"), WORD, False),
    "excel-to-pdf": (office_action("Excel.Application"), EXCEL, False),
    "ppt-to-pdf": (office_action("PowerPoint.Application"), POWERPOINT, False),
    "jpg-to-pdf": (images_to_pdf, IMAGES, True),
    "pdf-to-word": (pdf_to_word, PDF, False),
    "pdf-to-jpg": (pdf_to_jpg, PDF, False),
    "pdf-to-excel": (pdf_to_excel, PDF, False),
    "pdf-to-ppt": (pdf_to_ppt, PDF, False),
}


def save_uploads(uploads, tmp):
    """Grava os envios com nomes gerados aqui; o nome original só vira nome do download."""
    files = []
    for i, f in enumerate(uploads):
        original = Path(f.filename.replace('\\', '/').rsplit('/', 1)[-1])
        suffix = original.suffix.lower() if re.fullmatch(r'\.[a-z0-9]{1,5}', original.suffix.lower()) else ''
        path = tmp / f"entrada{i}{suffix}"
        f.save(path)
        files.append((path, original.stem or "arquivo"))
    return files


@app.route('/')
def index():
    # Extensões aceitas e se aceita vários arquivos, para os botões da interface
    tools = {action: {'accept': ','.join(accepted), 'multiple': multiple}
             for action, (_, accepted, multiple) in ACTIONS.items()}
    return render_template('index.html', tools=tools, bg_model_ready=bg_model_path().exists(), bg_model_mb=BG_MODEL_MB)

@app.route('/pages', methods=['POST'])
def page_info():
    """Número de páginas do PDF escolhido e, se pedidas, as miniaturas para organizar."""
    f = request.files.get('file')
    try:
        doc = open_pdf(f.read() if f else b'')
    except UserError as e:
        return {'erro': str(e)}, 400
    info = {'paginas': doc.page_count}
    if request.form.get('miniaturas'):
        # ponytail: todas as miniaturas numa resposta só; paginar se PDFs de centenas de páginas ficarem lentos
        info['miniaturas'] = [thumbnail(page) for page in doc]
    return info

def check_uploads(files, accepted, multiple):
    if not files:
        raise UserError("Escolha um arquivo.")
    if len(files) > 1 and not multiple:
        raise UserError("Esta ferramenta aceita um arquivo por vez.")
    for f in files:
        if Path(f.filename).suffix.lower() not in accepted:
            raise UserError(f"\"{f.filename}\" não é do tipo aceito aqui ({', '.join(accepted)}).")


def process_uploads(action, work):
    """Valida os arquivos enviados para a ferramenta, grava numa pasta temporária e devolve
    work(arquivos, pasta). A pasta é apagada ao fim, com o que houver dentro."""
    if action not in ACTIONS:
        raise UserError("Ferramenta desconhecida.")
    _, accepted, multiple = ACTIONS[action]
    files = [f for f in request.files.getlist('file') if f.filename]
    check_uploads(files, accepted, multiple)
    with tempfile.TemporaryDirectory(prefix='euamopdf-') as tmp:
        tmp = Path(tmp)
        saved = save_uploads(files, tmp)
        if any(path.stat().st_size == 0 for path, _ in saved):
            raise UserError("O arquivo enviado está vazio.")
        return work(saved, tmp)


@app.errorhandler(UserError)
def user_error(e):
    return str(e), 400, TEXT

@app.errorhandler(413)
def too_large(e):
    return f"Arquivo grande demais: o limite é {MAX_UPLOAD_MB} MB.", 413, TEXT

@app.errorhandler(Exception)
def unexpected_error(e):
    if isinstance(e, HTTPException):
        return e  # 404, 405...: a resposta padrão do Flask já serve
    app.logger.exception("Falha em %s", request.form.get('action'))
    return "Não foi possível processar o arquivo. Veja se ele abre normalmente em outro programa.", 500, TEXT

@app.route('/convert', methods=['POST'])
def handle_conversion():
    action = request.form.get('action')
    data, download_name, *message = process_uploads(
        action, lambda saved, tmp: ACTIONS[action][0](saved, request.form, tmp))
    response = send_file(io.BytesIO(data), as_attachment=True, download_name=download_name)
    if message:  # cabeçalhos HTTP só aceitam ASCII
        response.headers['X-Mensagem'] = quote(message[0])
    return response

def png_response(img):
    buf = io.BytesIO()
    img.save(buf, 'PNG')
    buf.seek(0)
    return send_file(buf, mimetype='image/png')

def single_image(saved):
    if len(saved) != 1:
        raise UserError("Envie uma imagem por vez.")
    return open_image(saved[0][0])[0]

@app.route('/background/mask', methods=['POST'])
def background_mask():
    """Máscara do objeto (PNG em tons de cinza, do tamanho da imagem), para a prévia editável."""
    return png_response(process_uploads('remove-background', lambda saved, tmp: subject_mask(single_image(saved))))

@app.route('/background/refine', methods=['POST'])
def background_refine():
    """Corrige a máscara a partir do traço pintado na prévia: restaura ou apaga a região marcada."""
    mode = request.form.get('modo')
    if mode not in ('restaurar', 'apagar'):
        raise UserError("Escolha entre restaurar e apagar.")

    smart = request.form.get('inteligente', '1') != '0'  # desligada: pincel exato

    def work(saved, tmp):
        img = single_image(saved)
        mask = load_mask(request.files.get('mascara'), img.size)
        painted = load_strokes(request.files.get('traco'), img.size, exact=not smart)
        if not smart:
            return paint_mask(mask, painted, mode == 'restaurar')
        return refine_mask(img, mask, painted, mode == 'restaurar')
    return png_response(process_uploads('remove-background', work))

@app.route('/metadata', methods=['POST'])
def current_metadata():
    """Metadados do PDF ou da imagem escolhida, para a interface preencher os campos."""
    f = request.files.get('file')
    data = f.read() if f else b''
    if f and not f.filename.lower().endswith('.pdf'):
        return read_image_metadata(data)
    return read_metadata(open_pdf(data)) | {'tipo': 'pdf'}

@app.route('/estimate', methods=['POST'])
def estimate_size():
    """Tamanho aproximado do resultado das ferramentas de compressão, com as opções escolhidas."""
    action = request.form.get('action')
    if action not in ESTIMATORS:
        raise UserError("Esta ferramenta não tem estimativa de tamanho.")
    before, after = process_uploads(action, lambda saved, tmp: ESTIMATORS[action](saved, request.form))
    return {'antes': before, 'depois': round(after)}

def remove_leftovers(max_age=86400):
    """Apaga as pastas de conversão que ficaram no temporário: fechar o terminal ou o app cair no
    meio de uma conversão não passa pela limpeza, e o documento ficava lá (no Windows, ninguém limpa).
    Só as de mais de um dia: as recentes podem ser de outra janela do app aberta."""
    limit = time.time() - max_age
    for folder in Path(tempfile.gettempdir()).glob('euamopdf-*'):
        try:
            if folder.is_dir() and folder.stat().st_mtime < limit:
                shutil.rmtree(folder, ignore_errors=True)
        except OSError:
            pass


def free_port(preferred=5000):
    """Usa a porta preferida se estiver livre; senão, qualquer porta livre."""
    with socket.socket() as s:
        try:
            s.bind(('127.0.0.1', preferred))
        except OSError:
            s.bind(('127.0.0.1', 0))
        return s.getsockname()[1]

if __name__ == '__main__':
    if hasattr(sys, '_MEIPASS'):
        # O PyInstaller aponta o LD_LIBRARY_PATH para as bibliotecas embutidas; o navegador e o
        # LibreOffice que o app abre precisam das do sistema. As nossas já foram carregadas.
        if 'LD_LIBRARY_PATH_ORIG' in os.environ:
            os.environ['LD_LIBRARY_PATH'] = os.environ['LD_LIBRARY_PATH_ORIG']
        else:
            os.environ.pop('LD_LIBRARY_PATH', None)
    remove_leftovers()
    port = free_port()
    print("O EuAmoPDF vive neste terminal: para encerrar, feche-o ou aperte Ctrl+C.")
    # O Timer aguarda 1 segundo para garantir que o servidor Flask já subiu
    Timer(1, webbrowser.open_new, [f"http://127.0.0.1:{port}/"]).start()
    app.run(host='127.0.0.1', port=port)
